import io
import csv
import logging
import uuid
import asyncio
import re
import zipfile
import xml.etree.ElementTree as ET
from typing import Optional, Literal
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Response, Query
from pydantic import BaseModel, Field
from i18n import t, message_catalog, set_override, reset_override, STRINGS
from db import db, get_settings, platform_db
from auth import get_current_admin, verify_password
from rates import get_rate
from pricing import price_for_product
from services import credit_deposit, reject_deposit, cancel_deposit, fmt_amount, now_iso, notify_product_created
from tgapi import download_telegram_file, send_message, send_photo_bytes, tg
from html import escape
from inventory import (
    validate_records,
    add_records,
    available_count,
    decrypt_items,
    encryption_status,
    InventoryError,
)
from storage import put_object, delete_object
from reporting import router as reports_router
from product_catalog import catalog_name as resolve_catalog_name
from catalog_routes import router as catalog_router
from inventory_admin import router as inventory_admin_router, align_records, parse_text_records, validate_schema
from product_artwork import artwork_urls, image_response, render_product_artwork

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/admin", dependencies=[Depends(get_current_admin)])
router.include_router(catalog_router)
router.include_router(inventory_admin_router)


# ============ STATS ============

@router.get("/stats")
async def stats():
    s = await get_settings()
    cutoff = s.get("stats_reset_at")

    def after_cutoff(match):
        if cutoff:
            return {**match, "created_at": {"$gte": cutoff}}
        return match

    async def sum_by(coll, match, field, currency):
        pipeline = [
            {"$match": {**after_cutoff(match), "currency": currency}},
            {"$group": {"_id": None, "t": {"$sum": "$" + field}}},
        ]
        res = await coll.aggregate(pipeline).to_list(1)
        return res[0]["t"] if res else 0

    dep_usd = await sum_by(db.deposits, {"status": "approved"}, "credited_amount", "USD")
    dep_idr = await sum_by(db.deposits, {"status": "approved"}, "credited_amount", "IDR")
    sales_usd = await sum_by(db.purchases, {}, "total", "USD")
    sales_idr = await sum_by(db.purchases, {}, "total", "IDR")
    circ = await db.bot_users.aggregate([
        {"$group": {"_id": None, "usd": {"$sum": "$balance_usd"}, "idr": {"$sum": "$balance_idr"}}}
    ]).to_list(1)
    pending = await db.deposits.count_documents(after_cutoff({"status": "pending"}))
    recent = await db.deposits.find(after_cutoff({})).sort("created_at", -1).to_list(8)
    rate = await get_rate()
    return {
        "total_deposit_usd": dep_usd,
        "total_deposit_idr": dep_idr,
        "total_sales_usd": sales_usd,
        "total_sales_idr": sales_idr,
        "circulating_usd": circ[0]["usd"] if circ else 0,
        "circulating_idr": circ[0]["idr"] if circ else 0,
        "pending_deposits": pending,
        "recent_deposits": recent,
        "rate": rate,
        "rate_mode": s.get("rate_mode"),
        "users_count": await db.bot_users.count_documents({}),
        "products_count": await db.products.count_documents({}),
        "stats_reset_at": cutoff,
    }


class ResetStatsBody(BaseModel):
    password: str


@router.post("/stats/reset")
async def reset_stats(body: ResetStatsBody, admin: dict = Depends(get_current_admin)):
    if not body.password:
        raise HTTPException(400, "Password wajib diisi.")
    stored = await platform_db().admins.find_one({"_id": admin["_id"]}, {"password_hash": 1})
    if not stored or not verify_password(body.password, stored.get("password_hash", "")):
        raise HTTPException(401, "Password admin salah.")

    reset_at = now_iso()
    await db.settings.update_one(
        {"_id": "main"},
        {"$set": {"stats_reset_at": reset_at}},
        upsert=True,
    )
    return {"ok": True, "stats_reset_at": reset_at}


# ============ PRODUCTS ============

def _normalized_product_kind(product: dict) -> str:
    kind = product.get("product_kind")
    if kind in {"digital", "service"}:
        return kind
    return "digital" if (
        product.get("delivery_type") == "inventory"
        or product.get("inventory_enabled")
    ) else "service"


def _is_inventory_product(product: dict) -> bool:
    return _normalized_product_kind(product) == "digital"


def _effective_admin_stock(product: dict, inventory_stock: int, marketing: int = 0) -> int | None:
    if not _is_inventory_product(product):
        return None
    mode = product.get("stock_mode", "auto")
    if mode == "manual" and product.get("manual_stock") is not None:
        return min(inventory_stock, max(0, int(product.get("manual_stock") or 0) - marketing))
    return inventory_stock


@router.get("/products")
async def list_products():
    products = await db.products.find().sort("created_at", -1).to_list(length=None)
    for product in products:
        product["catalog_name"] = resolve_catalog_name(product)
        kind = _normalized_product_kind(product)
        product["product_kind"] = kind
        if kind == "digital":
            actual = await available_count(product["_id"])
            product["inventory_stock"] = actual
            product["stock_mode"] = product.get("stock_mode", "auto")
            product["marketing_allocated"] = await db.inventory_items.count_documents({"product_id": product["_id"], "status": "marketing_allocated"})
            product["stock"] = _effective_admin_stock(product, actual, product["marketing_allocated"])
            product["inventory_enabled"] = True
        else:
            product["inventory_stock"] = 0
            product["stock"] = None
            product["inventory_enabled"] = False
        product.update(artwork_urls(product, admin=True))
    return products


@router.get("/products/artwork-preview")
async def product_artwork_preview(name: str = Query("Produk", max_length=300), catalog_name: str = Query("", max_length=80)):
    return Response(render_product_artwork(name, catalog_name), media_type="image/webp", headers={"Cache-Control": "private, max-age=60"})


@router.get("/products/{pid}/image")
async def admin_product_image(pid: str, catalog: bool = False):
    product = await db.products.find_one({"_id": pid})
    if not product:
        raise HTTPException(404, "Produk tidak ditemukan.")
    return await image_response(product, catalog)


def _validate_product_kind(value: str) -> str:
    value = (value or "digital").strip().lower()
    if value not in {"digital", "service"}:
        raise HTTPException(400, "Jenis product tidak valid.")
    return value


async def _store_product_image(product_id: str, image: Optional[UploadFile]):
    if not image or not image.filename:
        return None
    allowed = {"image/jpeg": "JPEG", "image/png": "PNG", "image/webp": "WEBP"}
    content_type = (image.content_type or "").lower()
    if content_type not in allowed:
        raise HTTPException(400, "Foto produk harus JPG, PNG, atau WEBP.")
    data = await image.read(5 * 1024 * 1024 + 1)
    if not data or len(data) > 5 * 1024 * 1024:
        raise HTTPException(400, "Ukuran foto maksimal 5 MB.")
    try:
        from PIL import Image, ImageOps
        with Image.open(io.BytesIO(data)) as decoded:
            expected_format = allowed[content_type]
            if decoded.format != expected_format:
                raise HTTPException(400, "Tipe foto tidak sesuai dengan isi file.")
            width, height = decoded.size
            if width < 32 or height < 32 or width > 6000 or height > 6000 or width * height > 25_000_000:
                raise HTTPException(400, "Dimensi foto harus antara 32 dan 6000 px, maksimal 25 megapiksel.")
            decoded.verify()
        with Image.open(io.BytesIO(data)) as source:
            source = ImageOps.exif_transpose(source)
            has_alpha = source.mode in {"RGBA", "LA"} or "transparency" in source.info
            source = source.convert("RGBA" if has_alpha else "RGB")
            resampling = getattr(Image, "Resampling", Image).LANCZOS
            fitted = ImageOps.contain(source, (1200, 1200), method=resampling)
            canvas = Image.new("RGBA", (1200, 1200), (255, 255, 255, 0))
            left = (1200 - fitted.width) // 2
            top = (1200 - fitted.height) // 2
            canvas.alpha_composite(fitted.convert("RGBA"), (left, top))
            output = io.BytesIO()
            canvas.save(output, format="WEBP", quality=90, method=6)
            data = output.getvalue()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(400, "File foto tidak valid atau rusak.") from exc
    path = f"product_images/{product_id}/{uuid.uuid4().hex}.webp"
    stored = await put_object(path, data, "image/webp")
    return {"path": stored["path"], "content_type": "image/webp"}


@router.post("/products")
async def create_product(
    name: str = Form(...),
    catalog_name: str = Form("", max_length=80),
    inventory_fields: str = Form("", max_length=4000),
    description: str = Form(""),
    price_usd: float = Form(...),
    price_idr: Optional[float] = Form(None),
    delivery_type: str = Form("link"),
    content: str = Form(""),
    active: bool = Form(True),
    minimum_purchase_qty: int = Form(1),
    stock: Optional[int] = Form(None),
    product_kind: str = Form("digital"),
    stock_mode: str = Form("auto"),
    inventory_mode: str = Form("table"),
    service_wait_minutes: Optional[int] = Form(None),
    service_message_template: str = Form(""),
    image: Optional[UploadFile] = File(None),
    remove_image: bool = Form(False),
    file: Optional[UploadFile] = File(None),
):
    product_kind = _validate_product_kind(product_kind)
    if minimum_purchase_qty < 1 or minimum_purchase_qty > 1000:
        raise HTTPException(400, "Minimum pembelian harus antara 1 dan 1000 pcs.")
    if stock_mode not in {"auto", "manual"}:
        raise HTTPException(400, "Mode stok tidak valid.")
    if inventory_mode not in {"table", "telegram_session"}:
        raise HTTPException(400, "Mode inventory tidak valid.")

    storage_path, original_filename = None, None
    product_id = str(uuid.uuid4())
    product_image = await _store_product_image(product_id, image)
    if product_kind == "service":
        wait_minutes = int(service_wait_minutes or 5)
        if wait_minutes not in {1, 5, 10, 25, 60}:
            raise HTTPException(400, "Waktu tunggu jasa harus 1, 5, 10, 25, atau 60 menit.")
        manual_stock = None
        stored_stock = None
        inventory_enabled = False
        stored_delivery = "service"
        stored_content = ""
        message_template = (service_message_template or "Jasa {product_name} sedang dalam antrean, harap tunggu {wait_minutes} untuk dapat menghubungi admin.").strip()
        if not message_template:
            raise HTTPException(400, "Pesan antrean jasa wajib diisi.")
    else:
        manual_stock = None if stock_mode == "auto" else max(0, int(stock or 0))
        stored_stock = manual_stock
        inventory_enabled = True
        stored_delivery = "inventory"
        stored_content = ""
        wait_minutes = None
        message_template = ""

    prod = {
        "_id": product_id,
        "name": name.strip(),
        "catalog_name": " ".join(catalog_name.split()),
        "description": description,
        "price_usd": price_usd,
        "price_idr": price_idr,
        "product_kind": product_kind,
        "delivery_type": stored_delivery,
        "content": stored_content,
        "storage_path": storage_path,
        "original_filename": original_filename,
        "service_wait_minutes": wait_minutes,
        "service_message_template": message_template,
        "active": active,
        "image_path": product_image["path"] if product_image else None,
        "image_content_type": product_image["content_type"] if product_image else None,
        "minimum_purchase_qty": minimum_purchase_qty,
        "stock": stored_stock,
        "stock_mode": stock_mode if product_kind == "digital" else "unlimited",
        "manual_stock": manual_stock,
        "inventory_enabled": inventory_enabled,
        "inventory_schema": (["Session File"] if inventory_mode == "telegram_session" else validate_schema(inventory_fields.split(",")) if inventory_fields.strip() else []) if product_kind == "digital" else [],
        "inventory_mode": inventory_mode if product_kind == "digital" else "table",
        "stock_notice_available": False if product_kind == "digital" else None,
        "stock_notice_count": 0,
        "created_at": now_iso(),
    }
    await db.products.insert_one(prod)
    try:
        await notify_product_created(prod)
    except Exception:
        logger.exception("Auto broadcast product baru gagal")
    return prod


@router.put("/products/{pid}")
async def update_product(
    pid: str,
    name: str = Form(...),
    catalog_name: Optional[str] = Form(None, max_length=80),
    inventory_fields: Optional[str] = Form(None, max_length=4000),
    description: str = Form(""),
    price_usd: float = Form(...),
    price_idr: Optional[float] = Form(None),
    delivery_type: str = Form("link"),
    content: str = Form(""),
    active: bool = Form(True),
    minimum_purchase_qty: int = Form(1),
    stock: Optional[int] = Form(None),
    product_kind: str = Form("digital"),
    stock_mode: str = Form("auto"),
    inventory_mode: str = Form("table"),
    service_wait_minutes: Optional[int] = Form(None),
    service_message_template: str = Form(""),
    image: Optional[UploadFile] = File(None),
    remove_image: bool = Form(False),
    file: Optional[UploadFile] = File(None),
    files: Optional[list[UploadFile]] = File(None),
):
    product = await db.products.find_one({"_id": pid})
    if not product:
        raise HTTPException(404, "Produk tidak ditemukan")

    product_kind = _validate_product_kind(product_kind)
    if minimum_purchase_qty < 1 or minimum_purchase_qty > 1000:
        raise HTTPException(400, "Minimum pembelian harus antara 1 dan 1000 pcs.")
    if stock_mode not in {"auto", "manual"}:
        raise HTTPException(400, "Mode stok tidak valid.")
    if inventory_mode not in {"table", "telegram_session"}:
        raise HTTPException(400, "Mode inventory tidak valid.")

    updates = {
        "name": name.strip(),
        "description": description,
        "price_usd": price_usd,
        "price_idr": price_idr,
        "product_kind": product_kind,
        "active": active,
        "minimum_purchase_qty": minimum_purchase_qty,
        "updated_at": now_iso(),
    }
    if (product_kind != _normalized_product_kind(product) or inventory_mode != product.get("inventory_mode", "table")) and await db.inventory_items.count_documents({"product_id": pid}):
        raise HTTPException(409, "Jenis/mode produk tidak dapat diganti saat masih memiliki inventory. Buat varian baru agar data pesanan tetap utuh.")
    if catalog_name is not None:
        updates["catalog_name"] = " ".join(catalog_name.split())
    if inventory_fields and product_kind == "digital":
        schema = validate_schema(inventory_fields.split(","))
        if product.get("inventory_schema"):
            align_records(schema, [], product["inventory_schema"])
        elif await db.inventory_items.count_documents({"product_id": pid}):
            raise HTTPException(409, "Produk sudah memiliki inventory. Schema tidak dapat diganti.")
        else:
            updates["inventory_schema"] = schema
    if image and remove_image:
        raise HTTPException(400, "Unggah foto baru atau hapus foto yang ada, bukan keduanya.")
    product_image = await _store_product_image(pid, image)
    if product_image:
        updates["image_path"] = product_image["path"]
        updates["image_content_type"] = product_image["content_type"]
    elif remove_image:
        updates["image_path"] = None
        updates["image_content_type"] = None

    if product_kind == "service":
        wait_minutes = int(service_wait_minutes or 5)
        if wait_minutes not in {1, 5, 10, 25, 60}:
            raise HTTPException(400, "Waktu tunggu jasa harus 1, 5, 10, 25, atau 60 menit.")
        message_template = (service_message_template or "Jasa {product_name} sedang dalam antrean, harap tunggu {wait_minutes} untuk dapat menghubungi admin.").strip()
        if not message_template:
            raise HTTPException(400, "Pesan antrean jasa wajib diisi.")
        updates.update({
            "delivery_type": "service",
            "content": "",
            "service_wait_minutes": wait_minutes,
            "service_message_template": message_template,
            "storage_path": None,
            "original_filename": None,
            "stock": None,
            "stock_mode": "unlimited",
            "manual_stock": None,
            "inventory_enabled": False,
            "inventory_mode": "table",
        })
    else:
        manual_stock = None if stock_mode == "auto" else max(0, int(stock or 0))
        updates.update({
            "delivery_type": "inventory",
            "content": "",
            "service_wait_minutes": None,
            "service_message_template": "",
            "storage_path": None,
            "original_filename": None,
            "stock": manual_stock,
            "stock_mode": stock_mode,
            "manual_stock": manual_stock,
            "inventory_enabled": True,
            "inventory_mode": inventory_mode,
        })

    changed = await db.products.update_one({"_id": pid}, {"$set": updates})
    if changed.matched_count != 1:
        if product_image:
            await delete_object(product_image["path"])
        raise HTTPException(409, "Produk berubah saat disimpan. Muat ulang lalu coba lagi.")
    old_image_path = product.get("image_path")
    if old_image_path and (product_image or remove_image):
        try:
            await delete_object(old_image_path)
        except Exception:
            logger.exception("Gagal menghapus file foto produk lama: %s", pid)
    if active and product_kind == "digital":
        from stock_monitor import schedule_stock_scan
        schedule_stock_scan(pid)
    result = await db.products.find_one({"_id": pid})
    if result and _is_inventory_product(result):
        result["inventory_stock"] = await available_count(pid)
        result["marketing_allocated"] = await db.inventory_items.count_documents({"product_id": pid, "status": "marketing_allocated"})
        result["stock"] = _effective_admin_stock(result, result["inventory_stock"], result["marketing_allocated"])
    return result


def _parse_num(v, idr: bool):
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("Rp", "").replace("$", "").strip()
    if idr:
        s = s.replace(".", "").replace(",", "")
    else:
        s = s.replace(",", "")
    return float(s)


@router.get("/products/import-template")
async def product_import_template(format: str = "xlsx"):
    """Generate the multi-sheet product + inventory workbook."""
    try:
        from openpyxl import Workbook
        from openpyxl.worksheet.datavalidation import DataValidation

        wb = Workbook()
        ws = wb.active
        ws.title = "product"
        headers = ["katalog", "product", "jenis (inventory/jasa)", "Harga"]
        ws.append(headers)
        if format in {"csv", "txt"}:
            out = io.StringIO()
            csv.writer(out, delimiter="|" if format == "txt" else ",").writerow(headers)
            return Response(out.getvalue().encode("utf-8-sig"), media_type="text/plain" if format == "txt" else "text/csv",
                headers={"Content-Disposition": f'attachment; filename="template-product.{format}"'})
        if format != "xlsx":
            raise HTTPException(400, "Pilih format xlsx, csv, atau txt.")
        for cell in ws[1]:
            cell.font = cell.font.copy(bold=True)

        type_validation = DataValidation(
            type="list",
            formula1='"inventory,jasa"',
            allow_blank=False,
        )
        wait_validation = DataValidation(
            type="list",
            formula1='"1,5,10,25,60"',
            allow_blank=True,
        )
        ws.add_data_validation(type_validation)
        type_validation.add("C2:C1000")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = "A1:D1000"

        widths = [30, 45, 28, 22]
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = width

        info = wb.create_sheet("Petunjuk")
        info_rows = [
            ["Bagian", "Aturan"],
            ["Header", "katalog | product | jenis (inventory/jasa) | Harga"],
            ["katalog", "Nama kelompok, contoh Claude Pro. Semua varian memakai katalog yang sama."],
            ["product", "Nama varian, contoh Claude Pro 1 Bulan. Nama yang sama memperbarui produk yang sudah ada."],
            ["jenis (inventory/jasa)", "Isi inventory untuk produk dengan data stok, atau jasa untuk layanan."],
            ["Harga", "Harga dalam Rupiah, contoh 100000. Isi angka tanpa Rp."],
            ["Inventory XLSX", "Opsional: sheet dengan nama sama persis seperti product. Header mengikuti data, contoh email,password."],
            ["Inventory CSV/TXT", "Upload terpisah lewat Kelola Inventory setelah memilih produk."],
            ["Kolom tambahan opsional", "Deskripsi, Minimum Pembelian, Kolom Inventory (nama field dipisahkan koma), Waktu Tunggu (menit), Pesan Jasa."],
            ["Template lama", "Header Nama Product, Jenis Product, Harga IDR/Harga USD tetap didukung."],
        ]
        for row in info_rows:
            info.append(row)
        info.freeze_panes = "A2"
        info.column_dimensions["A"].width = 32
        info.column_dimensions["B"].width = 100

        sample = wb.create_sheet("Contoh Digital")
        sample.append(["email", "password", "recovery", "2fa"])
        sample.append(["example@gmail.com", "password123", "recovery@example.com", "enabled"])
        sample.append(["akun2@gmail.com", "password456", "recovery2@example.com", "enabled"])
        for cell in sample[1]:
            cell.font = cell.font.copy(bold=True)

        output = io.BytesIO()
        wb.save(output)
        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="template-bulk-product.xlsx"'},
        )
    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("Gagal membuat template bulk product")
        raise HTTPException(500, f"Gagal membuat template Excel: {type(exc).__name__}")


@router.post("/products/import")
async def import_products(file: Optional[UploadFile] = File(None), content: str = Form("")):
    """Import products from the new multi-sheet workbook.

    Sheet product contains product metadata. Each digital product gets its
    own inventory sheet named exactly like the product. The first row of that
    sheet is the product-specific inventory schema.
    """
    from bulk_product_import import import_workbook

    data = await file.read() if file else content.encode("utf-8")
    filename = (file.filename or "").lower() if file else "manual.txt"
    flat = not filename.endswith(".xlsx")
    if flat:
        if not filename.endswith((".csv", ".txt")):
            raise HTTPException(400, "Gunakan XLSX, CSV, TXT, atau tempel data sesuai template.")
        from openpyxl import Workbook
        try:
            text_data = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise HTTPException(400, "File teks harus memakai encoding UTF-8.")
        wb = Workbook(); ws = wb.active; ws.title = "product"
        for row in csv.reader(io.StringIO(text_data), delimiter="|" if filename.endswith(".txt") else _detect_delimiter(text_data.splitlines()[0] if text_data.splitlines() else "")):
            ws.append(row)
        out = io.BytesIO(); wb.save(out); data = out.getvalue()

    try:
        return await import_workbook(data, metadata_only=flat)
    except HTTPException:
        raise
    except InventoryError:
        raise
    except Exception as exc:
        logger.exception("Bulk product import gagal")
        raise HTTPException(500, f"Bulk product import gagal ({type(exc).__name__}). Lihat log backend.")


def _stringify_cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _detect_delimiter(line: str):
    candidates = ["|", ",", ";", "\t"]
    return max(candidates, key=lambda d: line.count("\t" if d == "\\t" else d))


def _xlsx_col_index(cell_ref: str) -> int:
    letters = "".join(ch for ch in cell_ref if ch.isalpha()).upper()
    index = 0
    for char in letters:
        index = index * 26 + (ord(char) - ord("A") + 1)
    return max(0, index - 1)


def _xlsx_text(element):
    return "".join(element.itertext()) if element is not None else ""


def _read_xlsx_rows_fallback(data: bytes):
    """Read the first worksheet without parsing styles.

    Some XLSX files produced by spreadsheet/mobile apps contain a broken or
    missing styles.xml. openpyxl rejects those files even when the worksheet
    data itself is readable. XLSX is a ZIP of XML parts, so we can safely
    recover the cell values without loading the style information.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = set(archive.namelist())
            if "xl/workbook.xml" not in names:
                raise ValueError("workbook.xml tidak ditemukan")

            main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
            rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
            package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"

            workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
            rels = {}
            if "xl/_rels/workbook.xml.rels" in names:
                rels_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
                for rel in rels_root:
                    rel_id = rel.attrib.get("Id")
                    target = rel.attrib.get("Target")
                    if rel_id and target:
                        rels[rel_id] = target

            sheets = workbook_root.find(f"{{{main_ns}}}sheets")
            if sheets is None:
                raise ValueError("Tidak ada worksheet di workbook")

            first_sheet = next(iter(sheets), None)
            if first_sheet is None:
                raise ValueError("Worksheet kosong")

            rel_id = first_sheet.attrib.get(f"{{{rel_ns}}}id")
            target = rels.get(rel_id) if rel_id else None
            if target:
                target = target.lstrip("/")
                if not target.startswith("xl/"):
                    target = f"xl/{target}"
            else:
                target = "xl/worksheets/sheet1.xml"

            if target not in names:
                raise ValueError(f"Worksheet XML tidak ditemukan: {target}")

            shared_strings = []
            if "xl/sharedStrings.xml" in names:
                shared_root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
                for item in shared_root.findall(f"{{{main_ns}}}si"):
                    shared_strings.append(_xlsx_text(item))

            sheet_root = ET.fromstring(archive.read(target))
            sheet_data = sheet_root.find(f"{{{main_ns}}}sheetData")
            if sheet_data is None:
                return []

            parsed_rows = []
            for row_node in sheet_data.findall(f"{{{main_ns}}}row"):
                cells = {}
                max_col = -1
                for cell in row_node.findall(f"{{{main_ns}}}c"):
                    ref = cell.attrib.get("r", "")
                    col = _xlsx_col_index(ref)
                    if col < 0:
                        continue

                    cell_type = cell.attrib.get("t")
                    value = ""
                    if cell_type == "inlineStr":
                        inline = cell.find(f"{{{main_ns}}}is")
                        value = _xlsx_text(inline)
                    else:
                        value_node = cell.find(f"{{{main_ns}}}v")
                        raw = _xlsx_text(value_node)
                        if cell_type == "s" and raw:
                            try:
                                value = shared_strings[int(raw)]
                            except (ValueError, IndexError):
                                value = raw
                        elif cell_type == "b":
                            value = "TRUE" if raw == "1" else "FALSE"
                        else:
                            value = raw

                    cells[col] = value
                    max_col = max(max_col, col)

                parsed_rows.append([cells.get(i, "") for i in range(max_col + 1)])

            return parsed_rows
    except zipfile.BadZipFile as exc:
        raise ValueError("File bukan XLSX/ZIP yang valid") from exc
    except ET.ParseError as exc:
        raise ValueError("XML workbook/worksheet rusak") from exc


async def _parse_inventory_input(file: Optional[UploadFile], content: str, product: dict, files: Optional[list[UploadFile]] = None):
    schema = [str(x).strip() for x in (product.get("inventory_schema") or []) if str(x).strip()]
    records = []
    source_name = (file.filename or "").lower() if file else ""

    # Telegram Session mode supports selecting multiple real .session files.
    if files:
        if product.get("inventory_mode") != "telegram_session" and product.get("inventory_schema") != ["Session File"]:
            raise HTTPException(400, "Upload banyak file .session hanya tersedia untuk product Telegram Session.")
        import base64
        schema_from_file = ["Session File"]
        for upload in files:
            name = (upload.filename or "").strip()
            if not name.lower().endswith(".session"):
                raise HTTPException(400, f"File {name or '(tanpa nama)'} bukan file .session.")
            data = await upload.read()
            # Session inventory is identified by the uploaded .session file itself.
            # An empty file is still a valid inventory item; do not discard it.
            if len(data) > 10 * 1024 * 1024:
                raise HTTPException(400, f"File {name} melebihi batas 10 MB.")
            records.append({
                "Session File": name,
                "__file_name": name,
                "__file_data_b64": base64.b64encode(data).decode("ascii"),
            })
        # Every selected .session file is an inventory item, including
        # zero-byte test/placeholder files.
        if not records:
            raise HTTPException(400, "Tidak ada file .session yang dipilih.")
        return schema_from_file, records

    if file:
        data = await file.read()
        if source_name.endswith(".xlsx"):
            rows = None
            openpyxl_error = None
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
                rows = [list(row) for row in wb.active.iter_rows(values_only=True)]
            except Exception as exc:
                openpyxl_error = exc

            # Fallback for valid XLSX files whose styles.xml is malformed or
            # unsupported by openpyxl. This reads worksheet values directly.
            if rows is None:
                try:
                    rows = _read_xlsx_rows_fallback(data)
                except Exception as fallback_exc:
                    raise HTTPException(
                        400,
                        f"File XLSX inventory tidak valid: {fallback_exc}. "
                        f"openpyxl: {openpyxl_error}",
                    )

            rows = [row for row in rows if any(value not in (None, "") for value in row)]
            if not rows:
                raise HTTPException(400, "File inventory kosong.")

            # Spreadsheet apps often keep formatted/empty columns to the right
            # of the real table (e.g. A1 has "Session File" while B/C are
            # visually blank). Those trailing empty cells are not schema fields.
            first_row = list(rows[0])
            while first_row and first_row[-1] in (None, ""):
                first_row.pop()
            headers = [_stringify_cell(v) for v in first_row]
            if not headers or any(not header for header in headers) or len(set(headers)) != len(headers):
                raise HTTPException(400, "Header inventory tidak boleh kosong atau duplikat.")
            schema_from_file = headers
            for row in rows[1:]:
                values = list(row[:len(headers)])
                if len(values) < len(headers):
                    values.extend([""] * (len(headers) - len(values)))
                record = {headers[i]: _stringify_cell(values[i]) for i in range(len(headers))}
                if any(record.values()):
                    records.append(record)
        elif source_name.endswith(".csv"):
            text_data = data.decode("utf-8-sig", errors="ignore")
            rows = list(csv.reader(io.StringIO(text_data), delimiter=_detect_delimiter(text_data.splitlines()[0] if text_data.splitlines() else "")))
            rows = [row for row in rows if any(str(v or "").strip() for v in row)]
            if not rows:
                raise HTTPException(400, "File inventory kosong.")
            headers = [str(v or "").strip() for v in rows[0]]
            if not all(headers) or len(set(headers)) != len(headers):
                raise HTTPException(400, "Header inventory tidak boleh kosong atau duplikat.")
            schema_from_file = headers
            for row in rows[1:]:
                record = {headers[i]: str(row[i] if i < len(row) else "").strip() for i in range(len(headers))}
                if any(record.values()):
                    records.append(record)
        elif source_name.endswith(".txt"):
            schema_from_file, records = parse_text_records(data.decode("utf-8-sig"), schema)

        else:
            if product.get("inventory_mode") != "telegram_session" and product.get("inventory_schema") != ["Session File"]:
                raise HTTPException(
                    400,
                    "Produk ini memakai inventory tabel. Gunakan XLSX/CSV/TXT sesuai schema product. "
                    "File .session hanya tersedia untuk product dengan mode Telegram Session.",
                )
            if not source_name.endswith(".session"):
                raise HTTPException(400, "Mode Telegram Session hanya menerima file .session.")
            import base64
            if len(data) > 10 * 1024 * 1024:
                raise HTTPException(400, "Ukuran satu file inventory maksimal 10 MB.")
            filename = (file.filename or "inventory.session").strip() or "inventory.session"
            schema_from_file = ["Session File"]
            records.append({
                "Session File": filename,
                "__file_name": filename,
                "__file_data_b64": base64.b64encode(data).decode("ascii"),
            })
    else:
        schema_from_file, records = parse_text_records(content or "", schema)

    return align_records(schema_from_file, records, schema)


@router.post("/products/{pid}/inventory/validate")
async def validate_inventory(
    pid: str,
    content: str = Form(""),
    file: Optional[UploadFile] = File(None),
    files: list[UploadFile] = File([]),
):
    product = await db.products.find_one({"_id": pid})
    if not product:
        raise HTTPException(404, "Produk tidak ditemukan")
    if not _is_inventory_product(product):
        raise HTTPException(400, "Produk jasa tidak memiliki inventory.")

    schema, records = await _parse_inventory_input(file, content, product, files)
    try:
        check = await validate_records(pid, records, schema)
        # Validation is the point where a product's first inventory file
        # establishes its schema. Existing schemas are never overwritten.
        if not product.get("inventory_schema") and schema:
            await db.products.update_one(
                {"_id": pid},
                {"$set": {"inventory_schema": schema, "updated_at": now_iso()}},
            )
    except InventoryError:
        raise
    except Exception as exc:
        logger.exception("Validasi inventory gagal untuk produk %s", pid)
        raise InventoryError(f"Validasi inventory gagal ({type(exc).__name__}). Lihat log backend.") from exc
    return {
        "schema": schema,
        "valid_count": check["valid_count"],
        "duplicate_count": check["duplicate_count"],
        "preview": check["valid"][:5],
        "duplicates": check["duplicates"][:5],
    }


@router.post("/products/{pid}/inventory/import")
async def import_inventory(
    pid: str,
    content: str = Form(""),
    file: Optional[UploadFile] = File(None),
    files: list[UploadFile] = File([]),
):
    product = await db.products.find_one({"_id": pid})
    if not product:
        raise HTTPException(404, "Produk tidak ditemukan")
    if not _is_inventory_product(product):
        raise HTTPException(400, "Produk jasa tidak memiliki inventory.")

    schema, records = await _parse_inventory_input(file, content, product, files)
    try:
        result = await add_records(pid, records, schema)
    except InventoryError:
        raise
    except Exception as exc:
        logger.exception("Import inventory gagal untuk produk %s", pid)
        raise InventoryError(f"Import inventory gagal ({type(exc).__name__}). Lihat log backend.") from exc
    result.pop("valid", None)
    result.pop("duplicates", None)
    result["stock"] = await available_count(pid)
    result["schema"] = schema
    return result


@router.get("/inventory/status")
async def inventory_status():
    """Diagnosa konfigurasi enkripsi inventory (tanpa membocorkan key)."""
    return {"encryption": await encryption_status()}


class InventoryManualBody(BaseModel):
    data: dict[str, str]


@router.post("/products/{pid}/inventory/manual")
async def add_inventory_manual(pid: str, body: InventoryManualBody):
    product = await db.products.find_one({"_id": pid})
    if not product:
        raise HTTPException(404, "Produk tidak ditemukan")
    if not _is_inventory_product(product):
        raise HTTPException(400, "Produk jasa tidak memiliki inventory.")
    if product.get("inventory_mode") == "telegram_session" or product.get("inventory_schema") == ["Session File"]:
        raise HTTPException(400, "Inventory session harus diunggah sebagai file .session asli.")
    schema = [str(x).strip() for x in (product.get("inventory_schema") or []) if str(x).strip()]
    if not schema:
        raise HTTPException(400, "Tetapkan kolom inventory atau unggah template XLSX/CSV/TXT terlebih dahulu.")
    _, records = align_records(list(body.data), [body.data], schema)
    try:
        result = await add_records(pid, records, schema)
    except InventoryError:
        raise
    except Exception as exc:
        logger.exception("Input manual inventory gagal untuk produk %s", pid)
        raise InventoryError(f"Simpan data inventory gagal ({type(exc).__name__}). Lihat log backend.") from exc
    if result["created"] != 1:
        raise HTTPException(409, "Data inventory sudah ada atau tidak valid.")
    stock = await available_count(pid)
    return {
        "ok": True,
        "schema": schema,
        "stock": stock,
    }


@router.get("/products/{pid}/inventory")
async def inventory_list(pid: str, status: str = "available", offset: int = Query(0, ge=0), limit: int = Query(100, ge=1, le=500)):
    product = await db.products.find_one({"_id": pid})
    if not product:
        raise HTTPException(404, "Produk tidak ditemukan")
    if not _is_inventory_product(product):
        raise HTTPException(400, "Produk jasa tidak memiliki inventory.")

    q = {"product_id": pid}
    if status not in {"all", "available", "reserved", "sold", "marketing_allocated"}:
        raise HTTPException(400, "Status inventory tidak valid.")
    if status != "all":
        q["status"] = status
    cursor = db.inventory_items.find(q).sort([("created_at", -1), ("_id", 1)]).skip(offset).limit(limit)
    rows = []
    for item in await cursor.to_list(limit):
        row = {
            "_id": item["_id"],
            "status": item["status"],
            "order_id": item.get("order_id"),
            "user_tid": item.get("user_tid"),
            "created_at": item.get("created_at"),
            "sold_at": item.get("sold_at"),
        }
        if item.get("status") != "sold" and item.get("secret"):
            decrypted = decrypt_items([item])[0]
            if "__file_name" in decrypted or "__file_data_b64" in decrypted:
                row["item"] = {
                    "file": decrypted.get("__file_name") or decrypted.get("file") or "inventory.bin",
                    "is_file": True,
                }
            else:
                row["item"] = decrypted
        rows.append(row)
    return {
        "schema": product.get("inventory_schema") or [],
        "items": rows,
        "total": await db.inventory_items.count_documents(q),
        "offset": offset,
        "limit": limit,
        "available": await db.inventory_items.count_documents({"product_id": pid, "status": "available"}),
        "reserved": await db.inventory_items.count_documents({"product_id": pid, "status": "reserved"}),
        "sold": await db.inventory_items.count_documents({"product_id": pid, "status": "sold"}),
        "marketing_allocated": await db.inventory_items.count_documents({"product_id": pid, "status": "marketing_allocated"}),
    }


@router.get("/products/{pid}/inventory/summary")
async def inventory_summary(pid: str):
    product = await db.products.find_one({"_id": pid})
    if not product:
        raise HTTPException(404, "Produk tidak ditemukan")
    if not _is_inventory_product(product):
        raise HTTPException(400, "Produk jasa tidak memiliki inventory.")
    return {
        "available": await db.inventory_items.count_documents({"product_id": pid, "status": "available"}),
        "reserved": await db.inventory_items.count_documents({"product_id": pid, "status": "reserved"}),
        "sold": await db.inventory_items.count_documents({"product_id": pid, "status": "sold"}),
        "marketing_allocated": await db.inventory_items.count_documents({"product_id": pid, "status": "marketing_allocated"}),
    }


@router.delete("/products/{pid}/inventory/{item_id}")
async def delete_inventory_item(pid: str, item_id: str):
    product = await db.products.find_one({"_id": pid})
    if not product:
        raise HTTPException(404, "Produk tidak ditemukan")
    existing = await db.inventory_items.find_one({"_id": item_id, "product_id": pid}, {"marketing_audit": 1})
    if existing and existing.get("marketing_audit"):
        from marketing_campaigns import stock_audit
        await stock_audit(existing)
    result = await db.inventory_items.delete_one({
        "_id": item_id,
        "product_id": pid,
        "status": "available",
        "marketing_audit": (existing or {}).get("marketing_audit"),
    })
    if result.deleted_count != 1:
        raise HTTPException(400, "Item hanya bisa dihapus saat masih tersedia.")
    from stock_monitor import schedule_stock_scan
    schedule_stock_scan(pid)
    return {"ok": True, "stock": await available_count(pid)}


@router.patch("/products/{pid}/toggle")
async def toggle_product(pid: str):
    prod = await db.products.find_one({"_id": pid})
    if not prod:
        raise HTTPException(404, "Produk tidak ditemukan")
    await db.products.update_one({"_id": pid}, {"$set": {"active": not prod.get("active", True)}})
    return {"active": not prod.get("active", True)}


@router.delete("/products/{pid}")
async def delete_product(pid: str):
    product = await db.products.find_one({"_id": pid}, {"image_path": 1})
    if not product:
        raise HTTPException(404, "Produk tidak ditemukan")
    await db.products.delete_one({"_id": pid})
    if product.get("image_path"):
        try:
            await delete_object(product["image_path"])
        except Exception:
            logger.exception("Gagal menghapus file foto produk: %s", pid)
    return {"ok": True}


# ============ DEPOSITS ============

@router.get("/deposits")
async def list_deposits(status: Optional[str] = None):
    q = {"status": status} if status and status != "all" else {}
    return await db.deposits.find(q).sort("created_at", -1).to_list(500)


class DecisionBody(BaseModel):
    note: str = ""


@router.post("/deposits/{dep_id}/approve")
async def approve_deposit_api(dep_id: str, body: DecisionBody):
    dep = await db.deposits.find_one({"_id": dep_id})
    if not dep:
        raise HTTPException(404, "Deposit tidak ditemukan")
    if dep["status"] != "pending":
        raise HTTPException(400, f"Deposit sudah diproses ({dep['status']})")
    await credit_deposit(dep, note=body.note or "Disetujui via dashboard")
    return {"ok": True}


@router.post("/deposits/{dep_id}/reject")
async def reject_deposit_api(dep_id: str, body: DecisionBody):
    dep = await db.deposits.find_one({"_id": dep_id})
    if not dep:
        raise HTTPException(404, "Deposit tidak ditemukan")
    if dep["status"] != "pending":
        raise HTTPException(400, f"Deposit sudah diproses ({dep['status']})")
    await reject_deposit(dep, note=body.note)
    return {"ok": True}


@router.post("/deposits/{dep_id}/cancel")
async def cancel_deposit_api(dep_id: str):
    dep = await db.deposits.find_one({"_id": dep_id})
    if not dep:
        raise HTTPException(404, "Deposit tidak ditemukan")
    if dep["status"] != "approved":
        raise HTTPException(400, f"Hanya deposit disetujui yang bisa dibatalkan ({dep['status']})")
    await cancel_deposit(dep)
    return {"ok": True}


@router.get("/deposits/{dep_id}/proof")
async def deposit_proof(dep_id: str):
    dep = await db.deposits.find_one({"_id": dep_id})
    if not dep or not (dep.get("proof_file_id") or dep.get("proof_storage_path")):
        raise HTTPException(404, "Bukti tidak ditemukan")
    if dep.get("proof_storage_path"):
        data, content_type = await get_object(dep["proof_storage_path"])
        return Response(content=data, media_type=content_type)
    data = await download_telegram_file(dep["proof_file_id"])
    if not data:
        raise HTTPException(502, "Gagal mengambil file dari Telegram")
    return Response(content=data, media_type="image/jpeg")


# ============ USERS ============

@router.get("/users")
async def list_users():
    users = await db.bot_users.find().sort("created_at", -1).to_list(500)
    counts = {c["_id"]: c["n"] for c in await db.purchases.aggregate([{"$group": {"_id": "$user_tid", "n": {"$sum": 1}}}]).to_list(1000)}
    for u in users:
        u["_id"] = str(u["_id"])
        u["purchase_count"] = counts.get(u["telegram_id"], 0)
        u.pop("state", None)
        u.pop("state_data", None)
    return users


@router.get("/users/search")
async def search_users(
    search: str = "", status: str = "all", lang: str = "all",
    has_deposit: str = "all", has_order: str = "all",
    min_deposit: float | None = None, max_deposit: float | None = None,
    min_purchase: float | None = None, max_purchase: float | None = None,
    min_balance: float | None = None, max_balance: float | None = None,
    product_id: str = "", registered_from: str = "", registered_to: str = "",
):
    if status not in {"all", "active", "frozen"} or lang not in {"all", "id", "en"}:
        raise HTTPException(400, "Filter pengguna tidak valid.")
    if has_deposit not in {"all", "yes", "no"} or has_order not in {"all", "yes", "no"}:
        raise HTTPException(400, "Filter aktivitas tidak valid.")
    user_q = {}
    if status == "active": user_q["frozen"] = {"$ne": True}
    elif status == "frozen": user_q["frozen"] = True
    if lang != "all": user_q["lang"] = lang
    if registered_from or registered_to:
        user_q["created_at"] = {}
        if registered_from: user_q["created_at"]["$gte"] = registered_from
        if registered_to: user_q["created_at"]["$lt"] = registered_to
    if search.strip():
        raw = search.strip()
        username_term = raw[1:] if raw.startswith("@") else raw
        s = re.escape(username_term)
        clauses = [
            {"username": {"$regex": s, "$options": "i"}},
            {"first_name": {"$regex": re.escape(raw), "$options": "i"}},
            {"last_name": {"$regex": re.escape(raw), "$options": "i"}},
        ]
        if raw.isdigit():
            clauses.append({"telegram_id": int(raw)})
        user_q["$or"] = clauses
    # Bound the candidate set before loading customer activity; the UI searches
    # interactively and already exposes an empty-query overview separately.
    users = await db.bot_users.find(user_q).sort("created_at", -1).limit(500).to_list(500)
    tids = [u["telegram_id"] for u in users]
    if not tids: return []
    deposits = await db.deposits.find({"user_tid": {"$in": tids}, "status": "approved"}, {"user_tid": 1, "credited_amount": 1, "amount": 1}).to_list(10000)
    orders = await db.purchases.find({"user_tid": {"$in": tids}}, {"user_tid": 1, "total": 1, "status": 1, "items": 1}).to_list(20000)
    connected_tids = {
        account["tg_user_id"]
        async for account in db.tg_accounts.find(
            {"tg_user_id": {"$in": tids}, "status": "active", "session_encrypted": {"$type": "string"}},
            {"tg_user_id": 1},
        )
    }
    dep_by_user = {}
    for d in deposits: dep_by_user[d["user_tid"]] = dep_by_user.get(d["user_tid"], 0.0) + float(d.get("credited_amount") or d.get("amount") or 0)
    order_by_user, products_by_user = {}, {}
    for o in orders:
        tid = o["user_tid"]; order_by_user.setdefault(tid, {"count": 0, "spending": 0.0})
        if o.get("status") not in {"pending", "failed", "delivery_failed"}:
            order_by_user[tid]["count"] += 1; order_by_user[tid]["spending"] += float(o.get("total") or 0)
        for item in o.get("items", []):
            if item.get("product_id"): products_by_user.setdefault(tid, set()).add(item["product_id"])
    result = []
    for u in users:
        tid = u["telegram_id"]; dep_total = dep_by_user.get(tid, 0.0); stats = order_by_user.get(tid, {"count": 0, "spending": 0.0})
        balance = max(float(u.get("balance_usd") or 0), float(u.get("balance_idr") or 0))
        if has_deposit == "yes" and dep_total <= 0: continue
        if has_deposit == "no" and dep_total > 0: continue
        if has_order == "yes" and stats["count"] <= 0: continue
        if has_order == "no" and stats["count"] > 0: continue
        if min_deposit is not None and dep_total < min_deposit: continue
        if max_deposit is not None and dep_total > max_deposit: continue
        if min_purchase is not None and stats["spending"] < min_purchase: continue
        if max_purchase is not None and stats["spending"] > max_purchase: continue
        if min_balance is not None and balance < min_balance: continue
        if max_balance is not None and balance > max_balance: continue
        if product_id and product_id not in products_by_user.get(tid, set()): continue
        u["total_deposit"] = dep_total; u["order_count"] = stats["count"]; u["total_spending"] = stats["spending"]
        u["telegram_account_connected"] = tid in connected_tids
        u["purchased_product_ids"] = list(products_by_user.get(tid, set()))
        u.pop("state", None); u.pop("state_data", None)
        u.pop("deposit_credit_ids", None); u.pop("checkout_refund_ids", None); u.pop("deposit_debit_ids", None)
        result.append(u)
    return result

class AdjustBody(BaseModel):
    currency: Literal["USD", "IDR"]
    amount: float = Field(ge=-1_000_000_000, le=1_000_000_000, allow_inf_nan=False)
    reason: str = Field(default="", max_length=500)
    request_id: str | None = Field(default=None, min_length=8, max_length=100, pattern=r"^[A-Za-z0-9_-]+$")


@router.post("/users/{tid}/adjust")
async def adjust_balance(tid: int, body: AdjustBody, admin: dict = Depends(get_current_admin)):
    from balance_admin import Adjustment, adjust
    user = await db.bot_users.find_one({"telegram_id": tid})
    if not user:
        raise HTTPException(404, "Pengguna tidak ditemukan")
    if not __import__("math").isfinite(body.amount) or body.amount == 0:
        raise HTTPException(400, "Jumlah adjustment tidak valid")
    request = Adjustment(currency=body.currency, amount=abs(body.amount),
        direction="ADD" if body.amount > 0 else "SUBTRACT", reason=body.reason.strip() or "Penyesuaian admin (legacy)",
        request_id=body.request_id or str(uuid.uuid4()))
    result = await adjust("bot:" + str(user["_id"]), request, admin)
    if not result.get("replayed"):
        try:
            lang = user.get("lang") or "id"
            sign = "+" if body.amount > 0 else "-"
            await send_message(tid, t(lang, "adj_notice", amount=sign + fmt_amount(abs(body.amount), body.currency),
                reason=t(lang, "reason_label", r=body.reason) if body.reason else ""))
        except Exception:
            pass
    return result


class FreezeBody(BaseModel):
    reason: str = ""


@router.post("/users/{tid}/freeze")
async def freeze_user(tid: int, body: FreezeBody):
    user = await db.bot_users.find_one({"telegram_id": tid})
    if not user:
        raise HTTPException(404, "Pengguna tidak ditemukan")
    await db.bot_users.update_one({"telegram_id": tid}, {"$set": {"frozen": True, "frozen_reason": body.reason}})
    await db.freeze_log.insert_one({"_id": str(uuid.uuid4()), "user_tid": tid, "action": "freeze", "reason": body.reason, "created_at": now_iso()})
    lang = user.get("lang") or "id"
    reason = t(lang, "reason_label", r=body.reason) if body.reason else ""
    await send_message(tid, t(lang, "frozen_notice", reason=reason))
    return {"ok": True}


@router.post("/users/{tid}/unfreeze")
async def unfreeze_user(tid: int):
    user = await db.bot_users.find_one({"telegram_id": tid})
    if not user:
        raise HTTPException(404, "Pengguna tidak ditemukan")
    await db.bot_users.update_one({"telegram_id": tid}, {"$set": {"frozen": False, "frozen_reason": ""}})
    await db.freeze_log.insert_one({"_id": str(uuid.uuid4()), "user_tid": tid, "action": "unfreeze", "reason": "", "created_at": now_iso()})
    await send_message(tid, t(user.get("lang") or "id", "unfrozen_notice"))
    return {"ok": True}


# ============ SETTINGS ============

@router.get("/settings")
async def get_settings_api():
    s = await get_settings()
    s["current_rate"] = await get_rate()
    return s


class SettingsBody(BaseModel):
    crypto_addresses: dict
    bank_name: str = ""
    bank_account_number: str = ""
    bank_account_holder: str = ""
    qris_enabled: bool = False
    store_qris_enabled: bool = False
    gopay_qr_timeout_minutes: int = Field(default=5, ge=1, le=60)
    whatsapp_contact_number: str = Field(default="+628123456789", max_length=30)
    telegram_contact_target: str = Field(default="", max_length=255)
    bank_enabled: bool = True
    min_deposit_usd: float = 15.0
    min_deposit_idr: float = 50000.0
    admin_telegram_id: str = ""
    rate_mode: str = "auto"
    manual_rate: float = 16000.0
    max_deposit_usd: float = 100000.0
    max_deposit_idr: float = 100000000.0
    join_gate_enabled: bool = True
    join_gate_fail_open: bool = True
    required_channels: list[dict] = []
    auto_broadcast_new_product: bool = False
    transaction_success_channel_enabled: bool = False
    broadcast_auto_image_enabled: bool = False
    broadcast_channel_id: str = ""
    broadcast_group_ids: str = ""
    transaction_channel_ids: str = Field(default="", max_length=2000)
    recap_channel_ids: str = Field(default="", max_length=2000)
    stock_notifications_enabled: bool = True
    join_group_target: str = ""
    # Webshop broadcast banner + wajib join channel
    store_name: str = ""
    store_tagline: str = ""
    broadcast_enabled: bool = False
    broadcast_message: str = ""
    require_channel_join: bool = False
    channel_url: str = ""
    channel_name: str = ""


def _normalize_required_channel(channel: dict) -> dict:
    channel = dict(channel or {})
    channel_id = str(channel.get("channel_id") or channel.get("id") or "").strip()
    title = str(channel.get("title") or channel.get("name") or "").strip()
    username = str(channel.get("username") or "").strip()
    invite_link = str(channel.get("invite_link") or channel.get("join_link") or "").strip()
    return {
        "channel_id": channel_id,
        "title": title,
        "username": username,
        "invite_link": invite_link,
        "enabled": channel.get("enabled", True) is not False,
    }


@router.put("/settings")
async def update_settings(body: SettingsBody):
    data = body.model_dump()
    import re
    whatsapp_digits = re.sub(r"\D", "", str(data.get("whatsapp_contact_number") or ""))
    if whatsapp_digits and not 8 <= len(whatsapp_digits) <= 15:
        raise HTTPException(400, "Nomor WhatsApp harus berisi 8–15 digit.")
    if whatsapp_digits.startswith("0"):
        whatsapp_digits = "62" + whatsapp_digits[1:]
    data["whatsapp_contact_number"] = whatsapp_digits
    telegram_target = str(data.get("telegram_contact_target") or "").strip()
    if telegram_target:
        if telegram_target.startswith("@"): telegram_target = telegram_target[1:]
        if telegram_target.startswith(("https://t.me/", "http://t.me/", "https://telegram.me/", "http://telegram.me/")):
            data["telegram_contact_target"] = telegram_target
        elif re.fullmatch(r"[A-Za-z0-9_]{5,32}", telegram_target):
            data["telegram_contact_target"] = f"https://t.me/{telegram_target}"
        else:
            raise HTTPException(400, "Tujuan Telegram harus berupa username atau tautan t.me yang valid.")
    else:
        data["telegram_contact_target"] = ""
    channels = [_normalize_required_channel(ch) for ch in data.get("required_channels", [])]
    channels = [ch for ch in channels if ch["channel_id"]]
    if len(channels) > 3:
        raise HTTPException(400, "Maksimal 3 channel wajib join.")
    if data.get("join_gate_enabled") and not channels:
        data["join_gate_enabled"] = False
    data["required_channels"] = channels
    data["broadcast_channel_id"] = str(data.get("broadcast_channel_id") or "").strip()
    data["broadcast_group_ids"] = str(data.get("broadcast_group_ids") or "").strip()
    for field in ("broadcast_group_ids", "transaction_channel_ids", "recap_channel_ids"):
        values = list(dict.fromkeys(value.strip() for value in re.split(r"[,\n]+", data.get(field) or "") if value.strip()))
        if any(not re.fullmatch(r"(?:-\d+|@[A-Za-z0-9_]{5,32})", value) for value in values):
            raise HTTPException(400, f"{field}: gunakan ID -100… atau @username, satu per baris.")
        data[field] = "\n".join(values)
    data["join_group_target"] = str(data.get("join_group_target") or "").strip()
    await db.settings.update_one({"_id": "main"}, {"$set": data}, upsert=True)
    s = await get_settings()
    s["current_rate"] = await get_rate()
    return s


@router.post("/settings/join-gate/test")
async def test_required_channel(channel_id: str):
    channel_id = str(channel_id or "").strip()
    if not channel_id:
        raise HTTPException(400, "Channel ID wajib diisi.")
    try:
        chat = await tg("getChat", chat_id=channel_id)
        if not chat.get("ok"):
            raise HTTPException(400, f"Telegram tidak bisa mengakses channel: {chat.get('description', 'unknown error')}")
        me = await tg("getMe")
        bot_id = (me.get("result") or {}).get("id")
        member = await tg("getChatMember", chat_id=channel_id, user_id=bot_id)
        if not member.get("ok"):
            raise HTTPException(400, f"Gagal membaca status bot: {member.get('description', 'unknown error')}")
        status = (member.get("result") or {}).get("status")
        if status not in {"creator", "administrator"}:
            raise HTTPException(400, "Bot harus menjadi administrator di channel agar wajib join bisa bekerja.")
        return {
            "ok": True,
            "title": (chat.get("result") or {}).get("title") or "",
            "username": (chat.get("result") or {}).get("username") or "",
            "bot_status": status,
        }
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(400, f"Telegram error: {exc}")


# ============ ORDERS ============

@router.get("/orders")
async def list_orders(status: str = "all", search: str = "", limit: int = 200):
    q = {}
    if status != "all":
        q["status"] = status
    if search.strip():
        pattern = re.escape(search.strip())
        q["$or"] = [
            {"invoice_id": {"$regex": pattern, "$options": "i"}},
            {"username": {"$regex": pattern, "$options": "i"}},
            {"customer_email": {"$regex": pattern, "$options": "i"}},
        ]
        if search.strip().isdigit():
            q["$or"].append({"user_tid": int(search.strip())})
    return await db.purchases.find(q).sort("created_at", -1).limit(max(1, min(limit, 500))).to_list(max(1, min(limit, 500)))


@router.post("/orders/{oid}/complete")
async def complete_service_order(oid: str):
    order = await db.purchases.find_one({"_id": oid})
    if not order:
        raise HTTPException(404, "Order tidak ditemukan")
    if order.get("status") != "service_waiting":
        raise HTTPException(400, "Hanya pesanan jasa yang sedang menunggu yang bisa diselesaikan.")
    changed = await db.purchases.update_one(
        {"_id": oid, "status": "service_waiting"},
        {"$set": {"status": "delivered", "delivered_at": now_iso(), "delivery_error": None,
                  "service_completed_by_admin_at": now_iso()}},
    )
    if changed.modified_count != 1:
        raise HTTPException(409, "Status pesanan berubah. Muat ulang daftar pesanan.")
    email_sent = False
    if order.get("customer_email") or order.get("customer_id"):
        from storefront_routes import send_order_completion_email
        email_sent = await send_order_completion_email(oid)
    if order.get("user_tid") and not order.get("customer_id"):
        try:
            if order.get("bot2") or order.get("payment_scope") == "bot2":
                from bot2 import send2 as send_completed
            else:
                from bot import send_message as send_completed
            await send_completed(order["user_tid"], f"✅ Pesanan jasa <code>{escape(str(order.get('invoice_id') or oid))}</code> telah selesai.")
        except Exception:
            logger.exception("Could not send service completion for %s", oid)
    return {"ok": True, "email_sent": email_sent}


@router.post("/orders/{oid}/email/retry")
async def retry_order_email(oid: str):
    order = await db.purchases.find_one({"_id": oid, "status": "delivered"})
    if not order:
        raise HTTPException(404, "Pesanan selesai tidak ditemukan.")
    if order.get("delivery_email_status") != "failed":
        raise HTTPException(409, "Email tidak berstatus gagal atau sedang diproses.")
    await db.purchases.update_one({"_id": oid, "delivery_email_status": "failed"}, {"$unset": {
        "delivery_email_status": "", "delivery_email_error": "", "delivery_email_started_at": "",
    }})
    from storefront_routes import send_order_completion_email
    sent = await send_order_completion_email(oid)
    if not sent:
        raise HTTPException(502, "Email belum berhasil dikirim. Periksa konfigurasi SMTP dan log backend.")
    return {"ok": True, "email_sent": True}


@router.get("/orders/{oid}")
async def get_order(oid: str):
    order = await db.purchases.find_one({"_id": oid})
    if not order:
        raise HTTPException(404, "Order tidak ditemukan")
    return order


@router.post("/orders/{oid}/refund")
async def refund_order(oid: str):
    order = await db.purchases.find_one({"_id": oid})
    if not order:
        raise HTTPException(404, "Order tidak ditemukan")
    if order.get("status") not in {"failed", "delivery_failed"}:
        raise HTTPException(400, "Hanya order gagal yang bisa direfund otomatis.")

    field = "balance_usd" if order["currency"] == "USD" else "balance_idr"
    refund_key = f"refund:{oid}"
    if order.get("user_tid"):
        wallet = db.bot_users
        wallet_query = {"telegram_id": order["user_tid"]}
    elif order.get("customer_id"):
        customer = await db.store_customers.find_one({"_id": order["customer_id"]})
        if not customer:
            raise HTTPException(404, "Akun pelanggan tidak ditemukan.")
        if customer.get("telegram_id"):
            wallet = db.bot_users
            wallet_query = {"telegram_id": customer["telegram_id"]}
        else:
            wallet = db.store_customers
            wallet_query = {"_id": order["customer_id"]}
    else:
        raise HTTPException(400, "Akun tujuan refund tidak ditemukan.")
    result = await wallet.update_one(
        {**wallet_query, "refund_ids": {"$ne": refund_key}},
        {"$inc": {field: float(order["total"])}, "$addToSet": {"refund_ids": refund_key}},
    )
    if result.modified_count != 1:
        raise HTTPException(409, "Order sudah direfund.")

    await db.purchases.update_one(
        {"_id": oid},
        {"$set": {"status": "refunded", "refunded_at": now_iso(), "refund_reason": "Admin refund"}},
    )
    user = await db.bot_users.find_one({"telegram_id": order["user_tid"]}, {"lang": 1}) if order.get("user_tid") else None
    if order.get("user_tid") and user:
        await send_message(
            order["user_tid"],
            t(user.get("lang") or "id", "order_refunded", amount=fmt_amount(order["total"], order["currency"]), invoice=order["invoice_id"]),
        )
    return {"ok": True}


# ============ BOT MESSAGES ============

ALLOWED_PLACEHOLDERS = {
    "user_name", "username", "product_name", "quantity", "price", "balance",
    "invoice_id", "date", "order_id", "total", "currency", "amount",
    "network", "coin", "reason", "name", "lang", "cur", "stock", "short",
    "type", "desc", "bank", "account", "holder", "min", "r",
}
ALLOWED_HTML_TAGS = {"b", "strong", "i", "em", "u", "s", "code", "pre", "br", "a", "blockquote"}


def validate_bot_message(text: str, lang: str, key: str):
    placeholders = set(re.findall(r"\{([a-zA-Z_][a-zA-Z0-9_]*)\}", text))
    default_text = STRINGS.get(lang, {}).get(key) or STRINGS["id"].get(key, "")
    allowed = set(re.findall(r"\{([a-zA-Z_][a-zA-Z0-9_]*)\}", default_text))
    unknown = sorted(placeholders - allowed)
    tags = re.findall(r"</?([a-zA-Z][a-zA-Z0-9]*)", text)
    invalid_tags = sorted(set(tag.lower() for tag in tags) - ALLOWED_HTML_TAGS)
    if unknown:
        raise HTTPException(400, f"Placeholder tidak diizinkan untuk {lang}/{key}: {', '.join(unknown)}")
    if invalid_tags:
        raise HTTPException(400, f"HTML tag tidak diizinkan: {', '.join(invalid_tags)}")


@router.get("/messages")
async def list_messages():
    docs = []
    for lang in ("id", "en"):
        for key in message_catalog():
            override = await db.bot_messages.find_one({"lang": lang, "key": key})
            docs.append({
                "lang": lang,
                "key": key,
                "default": STRINGS[lang].get(key, STRINGS["id"].get(key, key)),
                "text": override.get("text") if override else STRINGS[lang].get(key, STRINGS["id"].get(key, key)),
                "custom": bool(override),
            })
    return docs


class MessageBody(BaseModel):
    text: str


@router.put("/messages/{lang}/{key}")
async def update_message(lang: str, key: str, body: MessageBody):
    if lang not in ("id", "en") or key not in message_catalog():
        raise HTTPException(404, "Message key tidak ditemukan")
    validate_bot_message(body.text, lang, key)

    current = await db.bot_messages.find_one({"lang": lang, "key": key})
    next_version = int((current or {}).get("version", 0)) + 1
    previous_text = (
        current.get("text") if current
        else STRINGS[lang].get(key, STRINGS["id"].get(key, key))
    )

    await db.bot_message_history.insert_one({
        "_id": str(uuid.uuid4()),
        "lang": lang,
        "key": key,
        "version": next_version,
        "text": body.text,
        "previous_text": previous_text,
        "created_at": now_iso(),
    })

    await db.bot_messages.update_one(
        {"lang": lang, "key": key},
        {
            "$set": {
                "text": body.text,
                "active": True,
                "version": next_version,
                "updated_at": now_iso(),
            }
        },
        upsert=True,
    )
    set_override(lang, key, body.text)
    return {
        "lang": lang,
        "key": key,
        "text": body.text,
        "version": next_version,
        "custom": True,
    }


@router.get("/messages/{lang}/{key}/history")
async def message_history(lang: str, key: str):
    if lang not in ("id", "en") or key not in message_catalog():
        raise HTTPException(404, "Message key tidak ditemukan")
    return await db.bot_message_history.find(
        {"lang": lang, "key": key}
    ).sort("version", -1).limit(50).to_list(50)


class RollbackMessageBody(BaseModel):
    version: int


@router.post("/messages/{lang}/{key}/rollback")
async def rollback_message(lang: str, key: str, body: RollbackMessageBody):
    if lang not in ("id", "en") or key not in message_catalog():
        raise HTTPException(404, "Message key tidak ditemukan")

    source = await db.bot_message_history.find_one({
        "lang": lang,
        "key": key,
        "version": body.version,
    })
    if not source:
        raise HTTPException(404, "Version tidak ditemukan")

    validate_bot_message(source["text"], lang, key)
    current = await db.bot_messages.find_one({"lang": lang, "key": key})
    next_version = int((current or {}).get("version", 0)) + 1

    await db.bot_message_history.insert_one({
        "_id": str(uuid.uuid4()),
        "lang": lang,
        "key": key,
        "version": next_version,
        "text": source["text"],
        "previous_text": current.get("text") if current else None,
        "rollback_from": body.version,
        "created_at": now_iso(),
    })

    await db.bot_messages.update_one(
        {"lang": lang, "key": key},
        {"$set": {
            "text": source["text"],
            "active": True,
            "version": next_version,
            "updated_at": now_iso(),
        }},
        upsert=True,
    )
    set_override(lang, key, source["text"])
    return {"ok": True, "version": next_version, "rollback_from": body.version}


class MessageTestBody(BaseModel):
    lang: str = "id"
    key: str
    text: str


@router.post("/messages/test")
async def test_message(body: MessageTestBody):
    if body.lang not in ("id", "en") or body.key not in message_catalog():
        raise HTTPException(404, "Message key tidak ditemukan")
    validate_bot_message(body.text, body.lang, body.key)
    settings = await get_settings()
    admin_id = str(settings.get("admin_telegram_id") or "")
    if not admin_id:
        raise HTTPException(400, "Telegram ID admin belum dikonfigurasi")

    sample = {
        "name": "Admin Preview",
        "user_name": "Admin Preview",
        "username": "@preview",
        "product_name": "Produk Contoh",
        "quantity": "2",
        "price": "Rp 20.000",
        "balance": "Rp 100.000",
        "invoice_id": "INV-20260921-0001",
        "order_id": "ORDER-PREVIEW",
        "total": "Rp 40.000",
        "currency": "IDR",
        "amount": "Rp 40.000",
        "payment_amount": "Rp 40.123",
        "network": "Polygon",
        "coin": "USDT",
        "reason": "Preview",
        "lang": "Indonesia",
        "cur": "IDR",
        "stock": "10",
        "short": "Rp 10.000",
        "type": "Inventory",
        "desc": "Preview",
        "bank": "BCA",
        "account": "123456",
        "holder": "Admin",
        "min": "Rp 50.000",
        "r": "Preview",
        "invoice": "INV-20260921-0001",
    }
    try:
        rendered = body.text.format(**sample)
    except KeyError as exc:
        raise HTTPException(400, f"Placeholder tidak bisa dirender: {exc}")
    result = await send_message(int(admin_id), rendered)
    if not result.get("ok"):
        raise HTTPException(502, result.get("description", "Telegram gagal mengirim test"))
    return {"ok": True}

@router.delete("/messages/{lang}/{key}")
async def reset_message(lang: str, key: str):
    await db.bot_messages.delete_one({"lang": lang, "key": key})
    reset_override(lang, key)
    return {"ok": True}


# ============ USERS / SEARCH ============

@router.get("/users/search-legacy")
async def search_users_legacy(
    search: str = "",
    status: str = "all",
    lang: str = "all",
    has_deposit: str = "all",
    has_order: str = "all",
    limit: int = 200,
):
    query = {}
    search = search.strip()
    if search:
        parts = []
        if search.isdigit():
            parts.append({"telegram_id": int(search)})
        parts.extend([
            {"username": {"$regex": re.escape(search), "$options": "i"}},
            {"first_name": {"$regex": re.escape(search), "$options": "i"}},
        ])
        query["$or"] = parts
    if status == "frozen":
        query["frozen"] = True
    elif status == "active":
        query["frozen"] = {"$ne": True}
    if lang in ("id", "en"):
        query["lang"] = lang

    if status not in {"all", "active", "frozen"} or lang not in {"all", "id", "en"}:
        raise HTTPException(400, "Filter pengguna tidak valid.")
    if has_deposit not in {"all", "yes", "no"} or has_order not in {"all", "yes", "no"}:
        raise HTTPException(400, "Filter aktivitas tidak valid.")
    bounded_limit = max(1, min(limit, 500))
    users = await db.bot_users.find(query).sort("created_at", -1).limit(bounded_limit).to_list(bounded_limit)
    if not users:
        return []
    tids = [user["telegram_id"] for user in users]

    deposit_map = {}
    dep_cur = db.deposits.find({"status": "approved", "user_tid": {"$in": tids}}, {"user_tid": 1, "amount": 1, "credited_amount": 1})
    async for dep in dep_cur:
        deposit_map.setdefault(dep["user_tid"], 0.0)
        deposit_map[dep["user_tid"]] += float(dep.get("credited_amount") or dep.get("amount") or 0)

    order_map = {}
    order_cur = db.purchases.find({"user_tid": {"$in": tids}}, {"user_tid": 1, "total": 1, "status": 1})
    async for order in order_cur:
        row = order_map.setdefault(order["user_tid"], {"count": 0, "spending": 0.0})
        if order.get("status") not in {"pending", "failed", "delivery_failed"}:
            row["count"] += 1
            row["spending"] += float(order.get("total") or 0)

    result = []
    for user in users:
        tid = user["telegram_id"]
        dep_total = deposit_map.get(tid, 0.0)
        stats = order_map.get(tid, {"count": 0, "spending": 0.0})
        if has_deposit == "yes" and dep_total <= 0:
            continue
        if has_deposit == "no" and dep_total > 0:
            continue
        if has_order == "yes" and stats["count"] <= 0:
            continue
        if has_order == "no" and stats["count"] > 0:
            continue
        user.pop("state", None)
        user.pop("state_data", None)
        user["total_deposit"] = dep_total
        user["order_count"] = stats["count"]
        user["total_spending"] = stats["spending"]
        result.append(user)

    return result


# ============ DISCOUNTS ============

class DiscountBody(BaseModel):
    name: str
    product_ids: list[str] = Field(default_factory=list)
    mode: str = "percent"
    value: float = 0.0
    min_qty: int = 1
    max_qty: Optional[int] = None
    fixed_currency: Optional[str] = None
    active: bool = True
    starts_at: Optional[str] = None
    ends_at: Optional[str] = None
    priority: int = 0


@router.get("/discounts")
async def list_discounts():
    return await db.discounts.find().sort([("priority", -1), ("created_at", -1)]).to_list(500)


def _validate_discount(body: DiscountBody):
    from datetime import datetime
    import math

    if not body.name.strip():
        raise HTTPException(400, "Nama discount wajib diisi.")
    if not math.isfinite(body.value):
        raise HTTPException(400, "Nilai discount harus berupa angka yang valid.")
    if body.mode not in {"percent", "fixed"}:
        raise HTTPException(400, "Mode discount harus percent atau fixed")
    if body.value <= 0:
        raise HTTPException(400, "Nilai discount harus lebih dari 0")
    if body.mode == "percent" and body.value > 100:
        raise HTTPException(400, "Persentase discount maksimal 100%")
    if body.min_qty < 1:
        raise HTTPException(400, "Quantity minimal 1")
    if body.priority < -1000 or body.priority > 1000:
        raise HTTPException(400, "Prioritas harus antara -1000 dan 1000.")
    if body.max_qty is not None and body.max_qty < body.min_qty:
        raise HTTPException(400, "Max quantity tidak boleh lebih kecil dari min quantity")
    if body.mode == "fixed" and body.fixed_currency not in {"IDR", "USD"}:
        raise HTTPException(400, "Nominal/unit harus memilih mata uang IDR atau USD")
    for field in ("starts_at", "ends_at"):
        value = getattr(body, field)
        if value:
            try:
                datetime.fromisoformat(value.replace("Z", "+00:00"))
            except (ValueError, AttributeError):
                raise HTTPException(400, f"Tanggal {field} tidak valid.")
    if body.starts_at and body.ends_at:
        start = datetime.fromisoformat(body.starts_at.replace("Z", "+00:00"))
        end = datetime.fromisoformat(body.ends_at.replace("Z", "+00:00"))
        if start.tzinfo is None:
            start = start.replace(tzinfo=__import__("datetime").timezone.utc)
        if end.tzinfo is None:
            end = end.replace(tzinfo=__import__("datetime").timezone.utc)
        if end <= start:
            raise HTTPException(400, "Tanggal berakhir harus setelah tanggal mulai.")
    if body.mode == "percent":
        body.fixed_currency = None


@router.post("/discounts")
async def create_discount(body: DiscountBody):
    _validate_discount(body)
    doc = {
        "_id": str(uuid.uuid4()),
        **body.model_dump(),
        "created_at": now_iso(),
        "updated_at": now_iso(),
    }
    await db.discounts.insert_one(doc)
    return doc


@router.put("/discounts/{did}")
async def update_discount(did: str, body: DiscountBody):
    _validate_discount(body)
    result = await db.discounts.update_one(
        {"_id": did},
        {"$set": {**body.model_dump(), "updated_at": now_iso()}},
    )
    if result.matched_count != 1:
        raise HTTPException(404, "Discount tidak ditemukan")
    return await db.discounts.find_one({"_id": did})


@router.patch("/discounts/{did}/toggle")
async def toggle_discount(did: str):
    doc = await db.discounts.find_one({"_id": did})
    if not doc:
        raise HTTPException(404, "Discount tidak ditemukan")
    active = not doc.get("active", True)
    await db.discounts.update_one({"_id": did}, {"$set": {"active": active, "updated_at": now_iso()}})
    return {"active": active}


@router.delete("/discounts/{did}")
async def delete_discount(did: str):
    result = await db.discounts.delete_one({"_id": did})
    if result.deleted_count != 1:
        raise HTTPException(404, "Discount tidak ditemukan")
    return {"ok": True}


# ============ BROADCAST ============

def _broadcast_user_query(query: dict | None = None):
    safe_query = dict(query or {})
    safe_query["silent_blocked"] = {"$ne": True}
    return safe_query


async def _broadcast_worker(
    broadcast_id: str,
    query: dict,
    text_body: str,
    photo_bytes: bytes | None,
    filename: str | None,
    button_text: str | None,
    button_url: str | None,
):
    success = failed = blocked = 0
    query = _broadcast_user_query(query)
    cursor = db.bot_users.find(query, {"telegram_id": 1})
    async for user in cursor:
        tid = user["telegram_id"]

        # Re-check immediately before sending. This covers a user being
        # silent-blocked after the broadcast job was queued.
        if await db.bot_users.find_one(
            {"telegram_id": tid, "silent_blocked": True},
            {"_id": 1},
        ):
            continue
        kb = None
        if button_text and button_url:
            kb = {"inline_keyboard": [[{"text": button_text, "url": button_url}]]}
        try:
            if photo_bytes:
                result = await send_photo_bytes(
                    tid,
                    photo_bytes,
                    filename or "broadcast.jpg",
                    caption=text_body,
                    kb=kb,
                )
            else:
                result = await send_message(tid, text_body, kb=kb)
            if result.get("ok"):
                success += 1
            else:
                failed += 1
                if result.get("error_code") == 403:
                    blocked += 1
                    await db.bot_users.update_one({"telegram_id": tid}, {"$set": {"blocked": True}})
        except Exception:
            failed += 1
        if (success + failed) % 20 == 0:
            await db.broadcasts.update_one(
                {"_id": broadcast_id},
                {"$set": {"success": success, "failed": failed, "blocked": blocked}},
            )
        await asyncio.sleep(0.08)

    await db.broadcasts.update_one(
        {"_id": broadcast_id},
        {
            "$set": {
                "status": "completed",
                "success": success,
                "failed": failed,
                "blocked": blocked,
                "finished_at": now_iso(),
            }
        },
    )


@router.post("/broadcasts/preview")
async def broadcast_preview(lang: str = Form("all"), search: str = Form(""), status: str = Form("all")):
    query = _broadcast_user_query()
    if lang in ("id", "en"): query["lang"] = lang
    if status == "active": query.update({"frozen": {"$ne": True}, "blocked": {"$ne": True}})
    elif status == "frozen": query["frozen"] = True
    if search.strip():
        s = re.escape(search.strip()); query["$or"] = [{"username": {"$regex": s, "$options": "i"}}, {"first_name": {"$regex": s, "$options": "i"}}]
    return {"total": await db.bot_users.count_documents(query)}


@router.post("/broadcasts")
async def create_broadcast(
    text: str = Form(...),
    lang: str = Form("all"),
    search: str = Form(""),
    status: str = Form("all"),
    button_text: str = Form(""),
    button_url: str = Form(""),
    photo: Optional[UploadFile] = File(None),
    product_id: str = Form(""),
    auto_image: bool = Form(False),
):
    if not text.strip():
        raise HTTPException(400, "Pesan broadcast kosong")
    if auto_image and not (await get_settings()).get("broadcast_auto_image_enabled", False):
        raise HTTPException(400, "Auto Generate Picture sedang OFF di Pengaturan.")
    query = _broadcast_user_query()
    if lang in ("id", "en"):
        query["lang"] = lang
    if status == "active":
        query["frozen"] = {"$ne": True}
        query["blocked"] = {"$ne": True}
    elif status == "frozen":
        query["frozen"] = True
    if search.strip():
        s = re.escape(search.strip())
        query["$or"] = [
            {"username": {"$regex": s, "$options": "i"}},
            {"first_name": {"$regex": s, "$options": "i"}},
        ]

    if button_url and not re.match(r"^(https?://|tg://)", button_url.strip(), re.I):
        raise HTTPException(400, "URL tombol harus http(s) atau tg://")
    total = await db.bot_users.count_documents(query)

    photo_bytes = None
    filename = None
    if photo:
        photo_bytes = await photo.read()
        filename = photo.filename
    elif auto_image and product_id:
        product = await db.products.find_one({"_id": product_id, "active": True})
        if not product:
            raise HTTPException(400, "Produk untuk gambar otomatis tidak ditemukan.")
        try:
            from broadcast_image import render_product_image
            stock = await available_count(product["_id"]) if _is_inventory_product(product) else None
            price = fmt_amount(
                product.get("price_idr") if product.get("price_idr") is not None else product.get("price_usd") or 0,
                "IDR" if product.get("price_idr") is not None else "USD",
            )
            photo_bytes = render_product_image(product.get("name") or "Product", price, stock, product.get("description") or "")
            filename = "product-broadcast.jpg"
        except Exception as exc:
            logger.exception("Generate broadcast image gagal")
            raise HTTPException(500, f"Gagal membuat gambar otomatis: {type(exc).__name__}")

    doc = {
        "_id": str(uuid.uuid4()),
        "text": text,
        "lang": lang,
        "status": "running",
        "success": 0,
        "failed": 0,
        "blocked": 0,
        "total": total,
        "created_at": now_iso(),
        "finished_at": None,
        "product_id": product_id or None,
        "auto_image": bool(auto_image),
    }
    await db.broadcasts.insert_one(doc)
    asyncio.create_task(_broadcast_worker(
        doc["_id"], query, text, photo_bytes, filename, button_text or None, button_url or None
    ))
    return doc


async def _broadcast_channel_id():
    from services import _broadcast_channel_id as configured_channel
    return await configured_channel()


async def _broadcast_channel_target():
    channel_id = await _broadcast_channel_id()
    if not channel_id:
        raise HTTPException(400, "Channel broadcast belum dikonfigurasi. Isi BROADCAST_CHANNEL_ID atau required channel di settings.")
    return channel_id


async def _build_product_broadcast():
    products = await db.products.find({"active": True}).sort("created_at", 1).to_list(500)
    digital_lines = []
    service_lines = []
    for p in products:
        kind = _normalized_product_kind(p)
        if kind == "digital":
            stock = await available_count(p["_id"])
            if stock > 0:
                digital_lines.append(f"▫️ <b>{escape(str(p.get('name') or 'Product'))}</b> → <b>{stock}</b>")
        else:
            service_lines.append(f"▫️ <b>{escape(str(p.get('name') or 'Jasa'))}</b>")
    lines = ["📢 <b>PRODUCT UPDATE — IDSE NETWORK CONNECT HUB</b>", ""]
    if digital_lines:
        lines += ["💻 <b>PRODUCT DIGITAL TERSEDIA • AVAILABLE STOCK</b>", *digital_lines, ""]
    if service_lines:
        lines += ["🛠️ <b>PRODUCT JASA TERSEDIA</b>", "✨ Tersedia sesuai permintaan • Unlimited", *service_lines, ""]
    if not digital_lines and not service_lines:
        lines += ["⚠️ <b>Saat ini belum ada product aktif yang tersedia.</b>", ""]
    lines += ["🚀 <b>Siap diproses • Cepat • Profesional</b>", "🛒 Silakan order melalui bot:", "🤖 @Idse_MarketBot"]
    return "\n".join(lines).strip()


async def _build_auto_broadcast(content: str):
    content = (content or "both").strip().lower()
    if content not in {"discount", "stock", "both"}:
        raise HTTPException(400, "Isi broadcast otomatis harus discount, stock, atau both.")

    products = await db.products.find({"active": True}).sort("created_at", 1).to_list(500)
    discount_lines = []
    stock_lines = []

    for p in products:
        name = escape(str(p.get("name") or "Product"))
        pricing = await price_for_product(p, "USD", 1)
        has_discount = float(pricing.get("discount_per_unit") or 0) > 0
        if has_discount:
            normal = fmt_amount(pricing["base_unit_price"], "USD")
            sale = fmt_amount(pricing["unit_price"], "USD")
            saved = fmt_amount(pricing["discount_per_unit"], "USD")
            discount_lines.append(f"▫️ <b>{name}</b> → {normal} ➜ <b>{sale}</b> (hemat {saved}/unit)")

        kind = _normalized_product_kind(p)
        if kind == "digital":
            stock = await available_count(p["_id"])
            if stock > 0:
                stock_lines.append(f"▫️ <b>{name}</b> → <b>{stock}</b>")
        else:
            stock_lines.append(f"▫️ <b>{name}</b> → <b>Unlimited</b>")

    lines = ["📢 <b>PROMO & STOCK UPDATE — IDSE NETWORK CONNECT HUB</b>", ""]
    if content in {"discount", "both"}:
        lines += ["🏷️ <b>HARGA DISKON TERSEDIA</b>"]
        lines += discount_lines or ["▫️ Belum ada product dengan harga diskon aktif."]
        lines.append("")
    if content in {"stock", "both"}:
        lines += ["📦 <b>STOCK TERSEDIA</b>"]
        lines += stock_lines or ["▫️ Saat ini belum ada stock tersedia."]
        lines.append("")
    lines += ["🚀 <b>Siap diproses • Cepat • Profesional</b>", "🛒 Silakan order melalui bot:", "🤖 @Idse_MarketBot"]
    return "\n".join(lines).strip()


@router.get("/broadcasts/channel-product-preview")
async def broadcast_channel_product_preview():
    return {"text": await _build_product_broadcast()}


@router.get("/broadcasts/auto-preview")
async def broadcast_auto_preview(content: str = "both"):
    return {"text": await _build_auto_broadcast(content)}


@router.post("/broadcasts/channel")
async def broadcast_channel(
    mode: str = Form(...),
    text: str = Form(""),
    target: str = Form("channel"),
    content: str = Form("both"),
    product_id: str = Form(""),
    auto_image: bool = Form(False),
):
    mode = (mode or "").strip().lower()
    target = (target or "channel").strip().lower()
    if auto_image and not (await get_settings()).get("broadcast_auto_image_enabled", False):
        raise HTTPException(400, "Auto Generate Picture sedang OFF di Pengaturan.")
    content = (content or "both").strip().lower()

    if mode == "manual":
        channel_id = await _broadcast_channel_target()
        if not text.strip():
            raise HTTPException(400, "Pesan manual wajib diisi.")
        body = text.strip()
        result = await send_message(channel_id, body)
        if not result.get("ok"):
            raise HTTPException(502, f"Telegram gagal mengirim ke channel: {result.get('description', 'unknown error')}")
        return {"ok": True, "mode": mode, "target": "channel", "channel_id": channel_id, "text": body}

    if mode == "product":
        channel_id = await _broadcast_channel_target()
        product = await db.products.find_one({"_id": product_id, "active": True}) if product_id else None
        if not product:
            raise HTTPException(400, "Pilih product terlebih dahulu.")
        stock = await available_count(product["_id"]) if _is_inventory_product(product) else None
        price = fmt_amount(
            product.get("price_idr") if product.get("price_idr") is not None else product.get("price_usd") or 0,
            "IDR" if product.get("price_idr") is not None else "USD",
        )
        body = (
            "🛒 <b>" + escape(str(product.get("name") or "Product")) + "</b>\n\n"
            + escape(str(product.get("description") or "").strip()) + "\n\n"
            + "Harga: <b>" + escape(price) + "</b>"
        )
        image = None
        if auto_image:
            try:
                from broadcast_image import render_product_image
                image = render_product_image(product.get("name") or "Product", price, stock, product.get("description") or "")
            except Exception as exc:
                raise HTTPException(500, f"Gagal membuat gambar otomatis: {type(exc).__name__}")
        result = await (send_photo_bytes(channel_id, image, "product-broadcast.jpg", caption=body) if image else send_message(channel_id, body))
        if not result.get("ok"):
            raise HTTPException(502, f"Telegram gagal mengirim ke channel: {result.get('description', 'unknown error')}")
        return {"ok": True, "mode": mode, "target": "channel", "channel_id": channel_id, "text": body, "auto_image": bool(image)}

    if mode == "products":
        channel_id = await _broadcast_channel_target()
        body = await _build_product_broadcast()
        result = await send_message(channel_id, body)
        if not result.get("ok"):
            raise HTTPException(502, f"Telegram gagal mengirim ke channel: {result.get('description', 'unknown error')}")
        return {"ok": True, "mode": mode, "target": "channel", "channel_id": channel_id, "text": body}

    if mode == "auto":
        if target not in {"channel", "users"}:
            raise HTTPException(400, "Target broadcast otomatis tidak valid.")
        body = await _build_auto_broadcast(content)

        if target == "channel":
            channel_id = await _broadcast_channel_target()
            result = await send_message(channel_id, body)
            if not result.get("ok"):
                raise HTTPException(502, f"Telegram gagal mengirim ke channel: {result.get('description', 'unknown error')}")
            return {
                "ok": True,
                "mode": mode,
                "target": "channel",
                "content": content,
                "channel_id": channel_id,
                "text": body,
            }

        query = _broadcast_user_query({"blocked": {"$ne": True}})
        total = await db.bot_users.count_documents(query)
        doc = {
            "_id": str(uuid.uuid4()),
            "text": body,
            "lang": "all",
            "status": "running",
            "success": 0,
            "failed": 0,
            "blocked": 0,
            "total": total,
            "created_at": now_iso(),
            "finished_at": None,
            "broadcast_type": "auto",
            "broadcast_target": "users",
            "broadcast_content": content,
        }
        photo_bytes = None
        filename = None
        if auto_image and product_id:
            product = await db.products.find_one({"_id": product_id, "active": True})
            if not product:
                raise HTTPException(400, "Produk untuk gambar otomatis tidak ditemukan.")
            from broadcast_image import render_product_image
            stock = await available_count(product["_id"]) if _is_inventory_product(product) else None
            price = fmt_amount(
                product.get("price_idr") if product.get("price_idr") is not None else product.get("price_usd") or 0,
                "IDR" if product.get("price_idr") is not None else "USD",
            )
            photo_bytes = render_product_image(product.get("name") or "Product", price, stock, product.get("description") or "")
            filename = "product-broadcast.jpg"
        await db.broadcasts.insert_one(doc)
        asyncio.create_task(_broadcast_worker(
            doc["_id"], query, body, photo_bytes, filename, None, None
        ))
        return {
            "ok": True,
            "mode": mode,
            "target": "users",
            "content": content,
            "queued": True,
            "total": total,
            "text": body,
        }

    raise HTTPException(400, "Mode broadcast channel tidak valid.")



@router.get("/broadcasts")
async def list_broadcasts():
    return await db.broadcasts.find({"kind": {"$ne": "marketing_campaign"}}).sort("created_at", -1).to_list(100)

@router.post("/products/import-test")
async def import_test():
    return {"ok": True, "message": "POST browser berhasil"}


@router.post("/products/upload-test")
async def upload_test(file: UploadFile = File(...)):
    data = await file.read()
    return {
        "ok": True,
        "filename": file.filename,
        "content_type": file.content_type,
        "size": len(data),
    }


# ============ REPORTS ============
# Mounted under /api/admin/reports with the same admin authentication dependency.
router.include_router(reports_router)
