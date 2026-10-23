"""Comprehensive IDSE Workbook Import and Apply Engine.

Implements end-to-end import:
- Catalogs creation/upsert
- Products creation/updating with full schema, structured warranties, combined descriptions,
  internal audit metadata, kind/delivery_type handling, and Fresh custom orders
- Stock & inventory parsing, encryption, deduplication/idempotency via fingerprint,
  and effective stock calculation
- Receipts recording with status COMPLETED and detailed stats
"""
from __future__ import annotations

import io
import json
import logging
import re
from copy import deepcopy
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from typing import Any, Dict, List, Optional, Tuple
from uuid import NAMESPACE_URL, uuid4, uuid5

import openpyxl

from idse_import_plan import record_version
from idse_inventory_reader import read_inventory
from idse_workbook_schema import inspect_workbook
from inventory import _canonical, _fernet, InventoryConfigError
from product_metadata import WarrantyMetadata

logger = logging.getLogger(__name__)


def parse_warranty(durasi_str: Optional[str], seller_garansi: Optional[str]) -> Dict[str, Any]:
    """Parse duration and seller warranty into structured WarrantyMetadata-compatible dict."""
    dur = (durasi_str or "").strip()
    if not dur or dur.lower() in ("tidak ada", "tidak ada garansi"):
        return {
            "enabled": False,
            "duration": None,
            "unit": None,
            "duration_value": None,
            "seller_warranty": seller_garansi or None,
            "coverage": None,
            "terms": None,
            "claim_process": None,
        }
    if "belum ditentukan" in dur.lower() or "belum" in dur.lower():
        return {
            "enabled": False,
            "duration": dur,
            "unit": "undetermined",
            "duration_value": None,
            "seller_warranty": seller_garansi or None,
            "coverage": None,
            "terms": None,
            "claim_process": None,
        }
    # Match pattern: e.g., '12 bulan', '2 jam', '25 hari'
    match = re.search(r"(\d+)\s*(bulan|hari|jam|tahun)", dur, re.IGNORECASE)
    if match:
        val = int(match.group(1))
        unit = match.group(2).lower()
        clean_dur = f"{val} {unit}"
        return {
            "enabled": True,
            "duration": clean_dur,
            "duration_value": val,
            "unit": unit,
            "seller_warranty": seller_garansi or None,
            "coverage": seller_garansi or None,
            "terms": None,
            "claim_process": "Klaim melalui ID order dan bukti pembelian/gangguan.",
        }
    return {
        "enabled": True,
        "duration": dur,
        "duration_value": None,
        "unit": "custom",
        "seller_warranty": seller_garansi or None,
        "coverage": seller_garansi or None,
        "terms": None,
        "claim_process": None,
    }


def join_full_description(
    summary_desc: Optional[str],
    desc_row: Dict[str, Any],
) -> str:
    """Join comprehensive description from Summary + Deskripsi & Aktivasi."""
    sections = []
    bot_summary = desc_row.get("bot_summary")
    features = desc_row.get("features")
    benefits = desc_row.get("benefits")
    activation = desc_row.get("activation")
    seller_warranty = desc_row.get("seller_warranty")
    official_sources = desc_row.get("official_sources")

    if bot_summary:
        sections.append(f"Ringkasan Bot:\n{str(bot_summary).strip()}")
    if features:
        sections.append(f"Fitur Utama:\n{str(features).strip()}")
    if benefits:
        sections.append(f"Benefit:\n{str(benefits).strip()}")
    if activation:
        sections.append(f"Cara Aktivasi:\n{str(activation).strip()}")
    if seller_warranty:
        sections.append(f"Garansi Seller:\n{str(seller_warranty).strip()}")
    if official_sources:
        sections.append(f"Sumber Resmi:\n{str(official_sources).strip()}")
    if summary_desc and str(summary_desc).strip():
        sections.append(f"Ringkasan Summary:\n{str(summary_desc).strip()}")

    return "\n\n".join(sections)


def parse_idse_workbook_for_apply(
    data: bytes,
    *,
    summary_precedence: bool = True,
) -> Dict[str, Any]:
    """Extract complete catalog, product, and inventory specifications ready for application."""
    report = inspect_workbook(data)
    if not report.get("structural_valid"):
        raise ValueError("Workbook memiliki kesalahan struktural: " + str(report.get("errors")))

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
    try:
        # 1. Parse Panduan & Mapping
        if "Panduan & Mapping" not in wb.sheetnames:
            raise ValueError("Sheet 'Panduan & Mapping' tidak ditemukan")
        mapping_sheet = wb["Panduan & Mapping"]
        mappings = []
        # Row 18 is header ['SKU', 'Katalog', 'Product', 'Jenis', 'Nama Sheet Stock']
        for row in mapping_sheet.iter_rows(max_col=5, values_only=True):
            sku, category, name, kind, target = row
            if isinstance(target, str) and target in wb.sheetnames and isinstance(sku, str):
                # Verify SKU matches cell B3 in detail target sheet if present
                if wb[target]["B3"].value == sku:
                    mappings.append((sku, category, name, kind, target))

        # 2. Parse Deskripsi & Aktivasi
        if "Deskripsi & Aktivasi" not in wb.sheetnames:
            raise ValueError("Sheet 'Deskripsi & Aktivasi' tidak ditemukan")
        desc_sheet = wb["Deskripsi & Aktivasi"]
        descriptions = {}
        for cells in desc_sheet.iter_rows(min_row=2, max_col=10):
            vals = [c.value for c in cells]
            if not vals[0]:
                continue
            sku_key = str(vals[0]).strip()
            descriptions[sku_key] = {
                "sku": sku_key,
                "catalog": vals[1],
                "product": vals[2],
                "features": vals[3],
                "benefits": vals[4],
                "activation": vals[5],
                "seller_warranty": vals[6],
                "risk_notes": vals[7],  # Internal audit only
                "official_sources": vals[8],
                "bot_summary": vals[9],
            }

        # 3. Parse Catatan Verifikasi (Internal audit notes)
        verifications = {}
        if "Catatan Verifikasi" in wb.sheetnames:
            verif_sheet = wb["Catatan Verifikasi"]
            for cells in verif_sheet.iter_rows(min_row=2, max_col=5):
                vals = [c.value for c in cells]
                if vals[0]:
                    sku_key = str(vals[0]).strip()
                    verifications[sku_key] = {
                        "sku": sku_key,
                        "product": vals[1],
                        "status": vals[2],
                        "importance": vals[3],
                        "steps": vals[4],
                    }

        # 4. Parse Summary and Detail Sheets
        summary_sheet = wb["Summary"]
        proposals = []
        catalogs = set()

        for number, row in enumerate(summary_sheet.iter_rows(min_row=2, max_col=8, values_only=True), 2):
            if not any(v is not None for v in row):
                continue
            _, category, name, stock_cell, kind, price, warranty_str, summary_desc = row
            category_clean = str(category or "").strip()
            if category_clean:
                catalogs.add(category_clean)

            # Locate SKU via mapping
            candidates = [m for m in mappings if m[1] == category and m[2] == name]
            if not candidates:
                candidates = [m for m in mappings if m[1] == category and (m[2] in name or name in m[2])]
            if len(candidates) != 1:
                # Stock formula link check
                if isinstance(stock_cell, str) and stock_cell.startswith("="):
                    match = re.fullmatch(r"='?([^'!]+)'?!B4", stock_cell)
                    if match:
                        candidates = [m for m in mappings if m[4] == match.group(1)]
            if len(candidates) != 1:
                raise ValueError(f"Mapping SKU ambigu atau tidak ditemukan untuk produk '{name}' di baris {number}")

            sku, _, _, mapped_kind, target_sheet = candidates[0]
            detail_sheet = wb[target_sheet]

            # Source kind precedence
            effective_kind = kind if summary_precedence else (detail_sheet["B5"].value or kind)
            # Price precedence
            amount = None
            effective_price = price if summary_precedence else (detail_sheet["B6"].value if detail_sheet["B6"].value is not None else price)
            if effective_price is not None:
                try:
                    val = Decimal(str(effective_price))
                    if val.is_finite() and val >= 0 and val == val.to_integral_value():
                        amount = int(val)
                except (InvalidOperation, ValueError):
                    amount = None

            desc_info = descriptions.get(sku, {})
            verif_info = verifications.get(sku, {})

            # Structured warranty
            warranty_dict = parse_warranty(warranty_str, desc_info.get("seller_warranty"))

            # Complete joined description
            full_description = join_full_description(summary_desc, desc_info)

            # Internal admin metadata
            internal_audit = {
                "risk_notes": desc_info.get("risk_notes"),
                "verification": verif_info,
                "official_sources": desc_info.get("official_sources"),
                "imported_at": datetime.now(timezone.utc).isoformat(),
            }

            # Inventory / Fresh handling
            custom_order_template = None
            fresh_status = "NONAKTIF"
            fresh_active = False
            inventory_items_to_create = []

            if effective_kind == "FRESH BY MADE":
                product_kind = "service"
                delivery_type = "service"
                b9_val = detail_sheet["B9"].value
                custom_order_template = str(b9_val).strip() if b9_val is not None else ""
                b14_val = detail_sheet["B14"].value
                fresh_status = str(b14_val or "").strip().upper()
                fresh_active = fresh_status == "AKTIF" and bool(custom_order_template)
            elif effective_kind == "Jasa":
                product_kind = "service"
                delivery_type = "service"
            else:  # Product
                product_kind = "digital"
                delivery_type = "inventory"
                # Read inventory from detail sheet starting row 9
                inv_result = read_inventory(detail_sheet)
                inventory_items_to_create = inv_result.units

            proposals.append({
                "sku": sku,
                "target_sheet": target_sheet,
                "summary_row": number,
                "name": name,
                "catalog_name": category_clean,
                "price_idr": amount,
                "source_kind": effective_kind,
                "product_kind": product_kind,
                "delivery_type": delivery_type,
                "warranty": warranty_dict,
                "description": full_description,
                "internal_audit": internal_audit,
                "custom_order_template": custom_order_template,
                "fresh_status": fresh_status,
                "fresh_active": fresh_active,
                "inventory_units": inventory_items_to_create,
                "summary_desc": summary_desc,
                "content": desc_info,
            })

        return {
            "catalogs": sorted(list(catalogs)),
            "proposals": proposals,
            "workbook_sha256": sha256(data).hexdigest(),
        }
    finally:
        wb.close()


async def apply_idse_import(
    database: Any,
    tenant_id: str,
    job_id: str,
    parsed_plan: Dict[str, Any],
    cipher: Any,
) -> Dict[str, Any]:
    """Execute full import apply with complete idempotency and audit trail."""
    catalogs = parsed_plan["catalogs"]
    proposals = parsed_plan["proposals"]
    now_iso = datetime.now(timezone.utc).isoformat()

    stats = {
        "catalogs_created": 0,
        "catalogs_existing": 0,
        "products_created": 0,
        "products_updated": 0,
        "products_skipped": 0,
        "inventory_inserted": 0,
        "inventory_skipped_duplicate": 0,
        "jasa_count": 0,
        "fresh_count": 0,
        "digital_count": 0,
        "errors": [],
    }

    # 1. Upsert product_catalogs
    for cat_name in catalogs:
        if not cat_name:
            continue
        cat_id = cat_name.lower().strip()
        existing_cat = await database.product_catalogs.find_one({"_id": cat_id})
        if existing_cat:
            stats["catalogs_existing"] += 1
        else:
            await database.product_catalogs.update_one(
                {"_id": cat_id},
                {"$setOnInsert": {"name": cat_name, "tenant_id": tenant_id, "created_at": now_iso}},
                upsert=True,
            )
            stats["catalogs_created"] += 1

    # 2. Upsert Products & Process Inventory
    for prop in proposals:
        sku = prop["sku"]
        kind = prop["source_kind"]
        pid = str(uuid5(NAMESPACE_URL, f"{tenant_id}:idse:{sku}"))
        receipt_id = str(uuid5(NAMESPACE_URL, f"{tenant_id}:{job_id}:{sku}"))

        if kind == "FRESH BY MADE":
            stats["fresh_count"] += 1
        elif kind == "Jasa":
            stats["jasa_count"] += 1
        else:
            stats["digital_count"] += 1

        existing = await database.products.find_one({"$or": [{"_id": pid}, {"sku": sku}, {"import_sku": sku}]})
        if existing and existing.get("tenant_id") not in (None, tenant_id):
            stats["errors"].append({"sku": sku, "error": "TENANT_MISMATCH"})
            continue

        actual_pid = existing["_id"] if existing else pid

        # Build product document fields
        # Note: checkout is blocked if price_idr is None (kept blank per requirements)
        # Store internal audit notes as internal metadata only
        product_doc_updates = {
            "name": prop["name"],
            "sku": sku,
            "import_sku": sku,
            "tenant_id": tenant_id,
            "catalog_name": prop["catalog_name"],
            "price_idr": prop["price_idr"],
            "price_usd": 0,
            "description": prop["description"],
            "warranty": prop["warranty"],
            "product_kind": prop["product_kind"],
            "delivery_type": prop["delivery_type"],
            "source_kind": kind,
            "internal_audit": prop["internal_audit"],
            "last_import_receipt": receipt_id,
            "last_import_job": job_id,
            "updated_at": now_iso,
        }

        # Rich catalog marketing content separated from delivery
        if prop.get("content"):
            product_doc_updates["catalog_details"] = prop["content"]

        if kind == "FRESH BY MADE":
            product_doc_updates["custom_order_template"] = prop["custom_order_template"]
            product_doc_updates["fresh_status"] = prop["fresh_status"]
            product_doc_updates["stock_mode"] = "manual"
            product_doc_updates["manual_stock"] = 1 if prop["fresh_active"] else 0
            product_doc_updates["stock"] = 1 if prop["fresh_active"] else 0
            product_doc_updates["active"] = prop["fresh_active"]
            product_doc_updates["inventory_enabled"] = False
        elif kind == "Jasa":
            product_doc_updates["stock_mode"] = "unlimited"
            product_doc_updates["stock"] = None
            product_doc_updates["active"] = True
            product_doc_updates["inventory_enabled"] = False
        else:  # Product (Digital)
            product_doc_updates["stock_mode"] = "auto"
            product_doc_updates["inventory_enabled"] = True
            # Default to active if price is set, review required if price is empty
            product_doc_updates["active"] = prop["price_idr"] is not None

        if existing is None:
            product_doc_updates["_id"] = actual_pid
            product_doc_updates["created_at"] = now_iso
            product_doc_updates["stock_notice_count"] = 0
            product_doc_updates["inventory_schema"] = ["email", "password"]
            await database.products.insert_one(product_doc_updates)
            stats["products_created"] += 1
        else:
            await database.products.update_one({"_id": actual_pid}, {"$set": product_doc_updates})
            stats["products_updated"] += 1

        # 3. Process Inventory Units (for Product kind)
        if kind == "Product":
            units = prop.get("inventory_units") or []
            schema_keys = set()
            for unit in units:
                raw_payload = unit["payload"]
                schema_keys.update(raw_payload.keys())
                canonical_str = _canonical(raw_payload)
                fingerprint = unit["fingerprint"]
                source_id = unit["source_id"]

                # Encrypt sensitive credentials
                encrypted_secret = cipher.encrypt(canonical_str.encode("utf-8")).decode("utf-8")

                # Idempotency check: check if already exists for this tenant & product
                inv_existing = await database.inventory_items.find_one({
                    "tenant_id": tenant_id,
                    "product_id": actual_pid,
                    "$or": [
                        {"fingerprint": fingerprint},
                        {"source_id": source_id},
                    ],
                })

                if inv_existing:
                    stats["inventory_skipped_duplicate"] += 1
                    continue

                inv_id = str(uuid4())
                inv_doc = {
                    "_id": inv_id,
                    "tenant_id": tenant_id,
                    "product_id": actual_pid,
                    "source_id": source_id,
                    "fingerprint": fingerprint,
                    "secret": encrypted_secret,
                    "status": "available",
                    "format": unit.get("format"),
                    "row_number": unit.get("row"),
                    "order_id": None,
                    "user_tid": None,
                    "reservation_id": None,
                    "reserved_at": None,
                    "sold_at": None,
                    "created_at": now_iso,
                    "movement_audit": [{
                        "event_id": str(uuid4()),
                        "action": "in",
                        "quantity": 1,
                        "at": now_iso,
                        "reference": f"excel_import:{job_id}",
                    }],
                }
                await database.inventory_items.insert_one(inv_doc)
                stats["inventory_inserted"] += 1

            # Recalculate available stock and update product document
            avail_count = await database.inventory_items.count_documents({
                "product_id": actual_pid,
                "status": "available",
            })
            await database.products.update_one(
                {"_id": actual_pid},
                {
                    "$set": {
                        "stock": avail_count,
                        "inventory_schema": sorted(list(schema_keys)) if schema_keys else (existing.get("inventory_schema") if existing else ["email", "password"]),
                    }
                },
            )

        # 4. Save per-product audit receipt
        await database.excel_import_receipts.update_one(
            {"_id": receipt_id},
            {
                "$set": {
                    "_id": receipt_id,
                    "tenant_id": tenant_id,
                    "job_id": job_id,
                    "sku": sku,
                    "product_id": actual_pid,
                    "status": "COMPLETED",
                    "applied_at": now_iso,
                }
            },
            upsert=True,
        )

    # 5. Record final import job summary receipt
    summary_receipt = {
        "_id": f"receipt:{job_id}",
        "tenant_id": tenant_id,
        "job_id": job_id,
        "status": "COMPLETED",
        "stats": stats,
        "applied_at": now_iso,
    }
    await database.excel_import_receipts.update_one(
        {"_id": summary_receipt["_id"]},
        {"$set": summary_receipt},
        upsert=True,
    )
    await database.excel_import_jobs.update_one(
        {"_id": job_id},
        {"$set": {"status": "COMPLETED", "completed_at": now_iso, "stats": stats}},
    )

    return {
        "status": "COMPLETED",
        "job_id": job_id,
        "stats": stats,
    }
