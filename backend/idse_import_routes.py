"""Tenant-admin workbook preflight and complete apply routes."""
import io
from typing import Optional
from uuid import uuid4

import openpyxl
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response

from idse_full_apply import apply_idse_import, parse_idse_workbook_for_apply
from idse_import_jobs import owned_plan, stage_plan
from idse_workbook_schema import (
    DESCRIPTION_HEADERS,
    SUMMARY_HEADERS,
    inspect_workbook,
)
from inventory import InventoryConfigError, _fernet
from tenant_context import get_tenant_context
from v2_commerce_routes import require_tenant_operator

router = APIRouter(
    prefix="/api/v2/import-excel",
    tags=["idse-import"],
    dependencies=[Depends(require_tenant_operator)],
)


@router.post("/preview")
async def preview(
    file: UploadFile = File(...),
    summary_precedence: bool = Form(False),
    context=Depends(get_tenant_context),
    actor=Depends(require_tenant_operator),
):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(400, "Unggah file .xlsx")
    data = await file.read(10 * 1024 * 1024 + 1)
    try:
        from idse_semantic_preview import semantic_preview

        result = semantic_preview(data, summary_precedence=summary_precedence)
    except Exception:
        # Parser exceptions may include XML/cell contents; never echo payload.
        raise HTTPException(400, "XLSX tidak valid atau melewati batas keamanan.") from None

    if result.get("products"):
        from idse_import_plan import compare_products

        result["database_diff"] = await compare_products(
            context.database, context.tenant_id, result["products"]
        )
        try:
            cipher = _fernet()
        except InventoryConfigError:
            raise HTTPException(
                503, "Penyimpanan preview terenkripsi belum dikonfigurasi."
            ) from None
        actor_id = actor.get("_id") or actor.get("user_id")
        if not actor_id:
            raise HTTPException(403, "Identitas admin tidak tersedia.")
        result["plan"] = await stage_plan(
            context.database, context.tenant_id, actor_id, data, result, cipher
        )

    # The template has 0 summary rows and no detail sheets/products.
    # apply_allowed is only True when products exist, structural_valid is True, and no blocking errors exist.
    has_blocking_errors = bool(result.get("errors"))
    has_products = bool(result.get("products"))
    apply_allowed = (
        result.get("structural_valid", False)
        and has_products
        and not has_blocking_errors
    )

    return {
        **result,
        "tenant_id": context.tenant_id,
        "phase": "metadata_preview",
        "apply_allowed": apply_allowed,
        "message": (
            "Preview metadata & inventory siap diterapkan."
            if apply_allowed
            else "Perbaiki kesalahan sebelum menerapkan import."
        ),
    }


@router.post("/apply")
async def apply(
    job_id: Optional[str] = Form(None),
    confirmation_token: Optional[str] = Form(None),
    file: Optional[UploadFile] = File(None),
    summary_precedence: bool = Form(True),
    context=Depends(get_tenant_context),
    actor=Depends(require_tenant_operator),
):
    """Apply the imported workbook to database.

    Accepts either (job_id, confirmation_token) from a previous preview,
    or direct (file, summary_precedence) upload.
    """
    try:
        cipher = _fernet()
    except InventoryConfigError:
        raise HTTPException(503, "Enkripsi belum dikonfigurasi.") from None

    actor_id = actor.get("_id") or actor.get("user_id")
    if not actor_id:
        raise HTTPException(403, "Identitas admin tidak tersedia.")

    data: Optional[bytes] = None
    active_job_id = job_id or f"direct_{uuid4().hex[:12]}"

    if job_id and confirmation_token:
        try:
            job, _ = await owned_plan(
                context.database,
                context.tenant_id,
                actor_id,
                job_id,
                confirmation_token,
                cipher,
            )
            data = cipher.decrypt(job["workbook_encrypted"].encode())
        except ValueError as exc:
            raise HTTPException(400, str(exc)) from exc
    elif file is not None:
        if not (file.filename or "").lower().endswith(".xlsx"):
            raise HTTPException(400, "Unggah file .xlsx")
        data = await file.read(10 * 1024 * 1024 + 1)
    else:
        raise HTTPException(
            400, "Berikan job_id & confirmation_token atau unggah file workbook .xlsx"
        )

    try:
        parsed_plan = parse_idse_workbook_for_apply(
            data, summary_precedence=summary_precedence
        )
    except Exception as exc:
        raise HTTPException(400, f"Gagal memproses workbook: {str(exc)}") from exc

    try:
        result = await apply_idse_import(
            context.database, context.tenant_id, active_job_id, parsed_plan, cipher
        )
        return {
            "ok": True,
            "tenant_id": context.tenant_id,
            **result,
            "message": "Import Excel berhasil diterapkan ke database.",
        }
    except Exception as exc:
        raise HTTPException(500, f"Gagal menerapkan import: {str(exc)}") from exc


@router.get("/template")
async def template(context=Depends(get_tenant_context)):
    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(SUMMARY_HEADERS)
    description = workbook.create_sheet("Deskripsi & Aktivasi")
    description.append(DESCRIPTION_HEADERS)
    mapping = workbook.create_sheet("Panduan & Mapping")
    mapping.append(["Template version", "1-draft"])
    mapping.append(["Status", "Template struktur siap pakai"])
    mapping.append(["Identitas", "SKU stabil wajib; jangan gunakan nama sebagai identitas"])
    output = io.BytesIO()
    workbook.save(output)
    return Response(
        output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="IDSE_Template_Structure_Draft.xlsx"',
            "Cache-Control": "no-store",
        },
    )
