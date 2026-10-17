"""No-write metadata preview. Summary precedence must be explicitly selected."""
import io
import re
from decimal import Decimal, InvalidOperation
import openpyxl
from idse_workbook_schema import inspect_workbook


def semantic_preview(data, *, summary_precedence=False):
    report = inspect_workbook(data)
    report['products'] = []
    if not report['structural_valid']:
        return report
    w = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        mappings = []
        if 'Panduan & Mapping' not in w.sheetnames:
            report['errors'].append({'code':'SKU_MAPPING_REQUIRED'})
            return report
        sheet = w['Panduan & Mapping']
        if sheet.max_row > 10000:
            report['errors'].append({'code':'MAPPING_LIMIT'})
            return report
        for row in sheet.iter_rows(max_col=5, values_only=True):
            sku, category, name, kind, target = row
            if isinstance(target,str) and target in w.sheetnames and isinstance(sku,str) and w[target]['B3'].value == sku:
                mappings.append((sku,category,name,kind,target))
        seen = set()
        for number, row in enumerate(w['Summary'].iter_rows(min_row=2,max_col=8,values_only=True),2):
            if not any(v is not None for v in row): continue
            _, category, name, stock, kind, price, warranty, description = row
            # Name/category locate an explicit mapping row only, never a DB product.
            # Stable identity is then cross-checked against detail B3.
            candidates = [m for m in mappings if m[1] == category and m[2] == name]
            if isinstance(stock,str) and stock.startswith('='):
                match = re.fullmatch(r"='?([^'!]+)'?!B4",stock)
                if match:
                    candidates = [m for m in candidates if m[4] == match.group(1)]
            if len(candidates) != 1:
                report['errors'].append({'sheet':'Summary','row':number,'code':'AMBIGUOUS_SKU_MAPPING'})
                continue
            sku, _, _, mapped_kind, target = candidates[0]
            issues = []
            if sku in seen: issues.append('DUPLICATE_SKU')
            seen.add(sku)
            if kind not in ('Product','Jasa','FRESH BY MADE'):issues.append('INVALID_KIND')
            detail = w[target]
            if (detail['B5'].value != kind or detail['B6'].value != price) and not summary_precedence:
                issues.append('PRECEDENCE_APPROVAL_REQUIRED')
            amount = None
            if price is not None:
                try:
                    value = Decimal(str(price))
                    if not value.is_finite() or value < 0 or value != value.to_integral_value(): raise ValueError()
                    amount = int(value)
                except (InvalidOperation,ValueError):issues.append('INVALID_PRICE')
            if amount is None:issues.append('PRICE_MISSING_CHECKOUT_BLOCKED')
            fresh_ready = False
            if kind == 'FRESH BY MADE':
                # Inventory-template rows are never interpreted as Fresh messages.
                fresh_ready = detail['A8'].value != 'Stock ID' and detail['B14'].value == 'AKTIF' and isinstance(detail['B9'].value,str) and bool(detail['B9'].value.strip()) and detail['B9'].data_type != 'f'
                if not fresh_ready:issues.append('FRESH_CONFIG_REQUIRED')
            inventory = {'valid_available':0,'skipped':0,'issues':[]}
            if kind == 'Product':
                from idse_inventory_reader import read_inventory
                inventory = read_inventory(detail).public()
            report['products'].append({'sku':sku,'source_sheet':target,'summary_row':number,
                'kind':kind,'price_idr':amount,'availability':int(fresh_ready) if kind=='FRESH BY MADE' else None,
                'action':'REVIEW','issues':issues,'inventory':inventory, 'stock_validation':'FAILED' if inventory['issues'] else 'PASS','publication':'REVIEW_REQUIRED'})
        report['apply_allowed'] = False
        return report
    finally:
        w.close()
