"""Reusable XLSX structural preflight. Never prints cell payloads or writes DB."""
import io
import zipfile
from hashlib import sha256
import openpyxl

SUMMARY_HEADERS = ('No', 'Katalog', 'Product', 'Stock', 'Jenis', 'Harga', 'Garansi', 'Deskripsi')
DESCRIPTION_HEADERS = ('SKU', 'Katalog', 'Product', 'Fitur Utama', 'Benefit', 'Cara Aktivasi', 'Garansi Seller', 'Catatan / Risiko', 'Sumber Resmi', 'Ringkasan Bot')


def inspect_workbook(data: bytes):
    if len(data) > 10 * 1024 * 1024:
        raise ValueError('XLSX exceeds 10 MB upload limit')
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(x.file_size for x in entries) > 60 * 1024 * 1024:
                raise ValueError('XLSX expanded size exceeds limit')
            if any('vbaproject' in x.filename.lower() or 'externallinks/' in x.filename.lower() for x in entries):
                raise ValueError('Macros and external workbook links are not accepted')
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except (zipfile.BadZipFile, KeyError, OSError):
        raise ValueError('Invalid XLSX container') from None
    errors = []
    try:
        for name, headers in [('Summary', SUMMARY_HEADERS), ('Deskripsi & Aktivasi', DESCRIPTION_HEADERS)]:
            if name not in workbook.sheetnames:
                errors.append({'sheet': name, 'code': 'MISSING_SHEET'})
                continue
            actual = tuple(workbook[name].cell(1, col).value for col in range(1,len(headers)+1))
            if actual != headers:
                errors.append({'sheet': name, 'row': 1, 'code': 'HEADER_MISMATCH', 'expected': list(headers)})
        count = 0
        if 'Summary' in workbook.sheetnames:
            sheet = workbook['Summary']
            if sheet.max_row > 10001 or sheet.max_column > 100:
                errors.append({'sheet': 'Summary', 'code': 'DIMENSION_LIMIT'})
            else:
                for row in sheet.iter_rows(min_row=2, max_col=8):
                    if any(cell.value is not None for cell in row):
                        count += 1
        return {'sha256': sha256(data).hexdigest(), 'sheet_count': len(workbook.sheetnames),
                'summary_rows': count, 'errors': errors, 'structural_valid': not errors,
                'apply_allowed': False}
    finally:
        workbook.close()
