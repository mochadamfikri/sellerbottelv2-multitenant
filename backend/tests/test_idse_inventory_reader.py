import openpyxl
from idse_inventory_reader import read_inventory, HEADERS


def test_formats_duplicates_and_status_are_safe():
    w=openpyxl.Workbook();s=w.active;s.title='P01'
    for c,v in enumerate(HEADERS,1):s.cell(8,c,v)
    rows=[['1','AVAILABLE','Email + Password','a@example.test','private'],
          ['2','AVAILABLE','Email + Password','a@example.test','private'],
          ['3','SOLD','Email + Password','sold@example.test','private'],
          ['4','RESERVED','Email + Password','held@example.test','private'],
          ['5','AVAILABLE','Email + Password','missing@example.test',None],
          ['6','AVAILABLE','Link',None,None,'https://example.test/redeem'],
          ['7','AVAILABLE','License',None,None,None,'LICENSE_PRIVATE'],
          ['8','AVAILABLE','Custom',None,None,None,None,'CUSTOM_PRIVATE']]
    for row in rows:s.append(row)
    result=read_inventory(s)
    assert len(result.units)==4 and result.skipped==3
    assert result.issues[0]['code']=='INCOMPLETE_CREDENTIAL_PAIR'
    assert 'private' not in str(result.public()).lower()
    assert 'LICENSE_PRIVATE' not in repr(result)


def test_formula_never_becomes_inventory():
    w=openpyxl.Workbook();s=w.active
    for c,v in enumerate(HEADERS,1):s.cell(8,c,v)
    s.append(['1','AVAILABLE','Custom',None,None,None,None,'=HYPERLINK("secret")'])
    r=read_inventory(s)
    assert not r.units and r.issues[0]['code']=='INVENTORY_FORMULA_FORBIDDEN'
