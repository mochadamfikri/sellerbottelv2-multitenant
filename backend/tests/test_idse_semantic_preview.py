import io
import openpyxl
from idse_semantic_preview import semantic_preview
from idse_workbook_schema import SUMMARY_HEADERS, DESCRIPTION_HEADERS


def workbook():
    w = openpyxl.Workbook(); s=w.active; s.title='Summary'
    s.append(SUMMARY_HEADERS)
    s.append([1,'Cat','Name','=P01!B4','FRESH BY MADE',95000,'Unknown','Long description'])
    m=w.create_sheet('Panduan & Mapping');m.append(['SKU1','Cat','Name','Product','P01'])
    d=w.create_sheet('P01');d['B3']='SKU1';d['B5']='Product';d['A8']='Stock ID'
    w.create_sheet('Deskripsi & Aktivasi').append(DESCRIPTION_HEADERS)
    out=io.BytesIO();w.save(out);return out.getvalue()


def test_precedence_never_silent_and_no_apply():
    r=semantic_preview(workbook())
    assert 'PRECEDENCE_APPROVAL_REQUIRED' in r['products'][0]['issues']
    approved=semantic_preview(workbook(),summary_precedence=True)
    p=approved['products'][0]
    assert p['sku']=='SKU1' and p['price_idr']==95000
    assert p['availability']==0 and 'FRESH_CONFIG_REQUIRED' in p['issues']
    assert not approved['apply_allowed']
