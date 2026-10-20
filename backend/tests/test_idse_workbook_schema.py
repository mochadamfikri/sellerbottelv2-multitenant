import io
import pytest
import openpyxl
from idse_workbook_schema import inspect_workbook, SUMMARY_HEADERS, DESCRIPTION_HEADERS


def sample(rows):
    w = openpyxl.Workbook()
    s = w.active
    s.title = 'Summary'
    s.append(SUMMARY_HEADERS)
    for i in range(rows):
        s.append([i+1,'Catalog','Product',0,'Product',100,'Unknown','Description'])
    s['J2'] = 'Internal extra summary'
    d = w.create_sheet('Deskripsi & Aktivasi')
    d.append(DESCRIPTION_HEADERS)
    out = io.BytesIO()
    w.save(out)
    return out.getvalue()


@pytest.mark.parametrize('rows', [1,21,35])
def test_reusable_counts_and_no_apply(rows):
    result = inspect_workbook(sample(rows))
    assert result['structural_valid']
    assert result['summary_rows'] == rows
    assert not result['apply_allowed']


def test_invalid_container():
    with pytest.raises(ValueError, match='Invalid XLSX'):
        inspect_workbook(b'not a spreadsheet')


def test_bad_headers_do_not_echo_private_values():
    w = openpyxl.load_workbook(io.BytesIO(sample(1)))
    w['Summary']['A1'] = 'PRIVATE_PAYLOAD'
    out = io.BytesIO(); w.save(out)
    result = inspect_workbook(out.getvalue())
    assert not result['structural_valid']
    assert 'PRIVATE_PAYLOAD' not in str(result)
