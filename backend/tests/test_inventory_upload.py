"""Tes upload/import inventory (halaman Product dan halaman Kelola Inventory).

Tes ini berjalan offline: MongoDB diganti ``mongomock-motor`` dan autentikasi admin
di-override, jadi aman dijalankan di CI tanpa server apa pun.
"""
import asyncio
import io
import os
import sys
import tempfile
import zipfile
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "sellerbottel_test")
os.environ.setdefault("JWT_SECRET", "test-secret-test-secret-test-secret")

import motor.motor_asyncio as _motor
from mongomock_motor import AsyncMongoMockClient, AsyncMongoMockCollection

_motor.AsyncIOMotorClient = AsyncMongoMockClient  # harus sebelum `db` di-import

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import httpx  # noqa: E402
import openpyxl  # noqa: E402
import pytest  # noqa: E402
from cryptography.fernet import Fernet  # noqa: E402
from fastapi import APIRouter, FastAPI  # noqa: E402

import admin_routes  # noqa: E402
import inventory  # noqa: E402
import inventory_transform  # noqa: E402
from auth import get_current_admin  # noqa: E402
from db import db  # noqa: E402
from error_handlers import register_error_handlers  # noqa: E402

PID = "prod-yt"
HEADER = ("email", "password", "recovery_email", "2FA")
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

boom = APIRouter()


@boom.get("/api/boom")
async def _boom():
    raise RuntimeError("kesalahan tak terduga")


app = FastAPI()
app.dependency_overrides[get_current_admin] = lambda: {"email": "admin@example.com"}
app.include_router(admin_routes.router)
app.include_router(inventory_transform.router)
app.include_router(boom)
register_error_handlers(app)


def run(coro):
    return asyncio.run(coro)


def request(method, path, **kwargs):
    async def _go():
        transport = httpx.ASGITransport(app=app, raise_app_exceptions=False)
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
            return await client.request(method, path, **kwargs)

    return run(_go())


def make_xlsx(rows, header=HEADER):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(header))
    for row in rows:
        ws.append(list(row))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def sample_rows(n=13, prefix="u"):
    return [(f"{prefix}{i}@x.com", f"pw{i}", f"r{i}@x.com", f"2fa{i}") for i in range(n)]


def upload(step, data, filename="upload inventory.xlsx", mime=XLSX_MIME, pid=PID):
    files = {"file": (filename, data, mime)}
    return request("POST", f"/api/admin/products/{pid}/inventory/{step}", files=files, data={"content": ""})


def test_empty_telegram_session_upload_from_both_pages():
    """An empty .session is inventory whether sent as file or files."""
    from fastapi import UploadFile

    async def check():
        product = {"inventory_mode": "telegram_session", "inventory_schema": ["Session File"]}
        for field in ("file", "files"):
            upload = UploadFile(file=tempfile.SpooledTemporaryFile(), filename=f"{field}.session")
            args = (upload, "", product) if field == "file" else (None, "", product)
            kwargs = {} if field == "file" else {"files": [upload]}
            schema, records = await admin_routes._parse_inventory_input(*args, **kwargs)
            assert schema == ["Session File"]
            assert len(records) == 1
            assert records[0]["__file_data_b64"] == ""
            checked = await inventory.validate_records(PID, records, schema)
            assert checked["valid_count"] == 1
            imported = await inventory.add_records(PID, records, schema)
            assert imported["created"] == 1

    run(check())


@pytest.fixture(autouse=True)
def clean_state(monkeypatch):
    monkeypatch.setenv("INVENTORY_ENCRYPTION_KEY", Fernet.generate_key().decode())

    async def reset():
        await db.inventory_items.delete_many({})
        await db.inventory_items.create_index([("product_id", 1), ("fingerprint", 1)], unique=True)
        await db.products.delete_many({})
        await db.products.insert_one({
            "_id": PID,
            "name": "YT Prem 3 bulan",
            "product_kind": "digital",
            "delivery_type": "inventory",
            "inventory_enabled": True,
            "inventory_schema": [],
        })

    run(reset())
    yield


def test_xlsx_validate_then_import_succeeds():
    data = make_xlsx(sample_rows())
    validated = upload("validate", data)
    assert validated.status_code == 200, validated.text
    assert validated.json()["valid_count"] == 13
    assert validated.json()["duplicate_count"] == 0

    imported = upload("import", data)
    assert imported.status_code == 200, imported.text
    body = imported.json()
    assert body["created"] == 13
    assert body["skipped"] == 0
    assert body["stock"] == 13
    assert body["schema"] == list(HEADER)


def test_reimporting_same_file_reports_duplicates_not_error():
    data = make_xlsx(sample_rows())
    assert upload("import", data).status_code == 200
    again = upload("import", data)
    assert again.status_code == 200, again.text
    assert again.json()["created"] == 0
    assert again.json()["skipped"] == 13
    assert again.json()["stock"] == 13


def test_import_without_key_returns_clear_error_not_blank_500():
    """Regresi bug utama: validasi lolos tetapi import 500 tanpa pesan."""
    os.environ["INVENTORY_ENCRYPTION_KEY"] = ""
    data = make_xlsx(sample_rows())
    assert upload("validate", data).status_code == 200
    res = upload("import", data)
    assert res.status_code == 503
    assert "INVENTORY_ENCRYPTION_KEY" in res.json()["detail"]
    assert run(db.inventory_items.count_documents({})) == 0


def test_import_with_invalid_key_returns_clear_error():
    os.environ["INVENTORY_ENCRYPTION_KEY"] = "bukan-fernet-key"
    res = upload("import", make_xlsx(sample_rows()))
    assert res.status_code == 503
    assert "tidak valid" in res.json()["detail"]


def test_status_endpoint_reports_key_state():
    ok = request("GET", "/api/admin/inventory/status").json()["encryption"]
    assert ok["configured"] and ok["valid"] and ok["data_readable"] is None

    assert upload("import", make_xlsx(sample_rows(3))).status_code == 200
    readable = request("GET", "/api/admin/inventory/status").json()["encryption"]
    assert readable["data_readable"] is True and readable["message"] is None

    os.environ["INVENTORY_ENCRYPTION_KEY"] = Fernet.generate_key().decode()
    mismatch = request("GET", "/api/admin/inventory/status").json()["encryption"]
    assert mismatch["valid"] is True
    assert mismatch["data_readable"] is False
    assert "tidak cocok" in mismatch["message"]

    os.environ["INVENTORY_ENCRYPTION_KEY"] = ""
    missing = request("GET", "/api/admin/inventory/status").json()["encryption"]
    assert missing["configured"] is False and missing["valid"] is False


def test_status_never_leaks_the_key():
    key = os.environ["INVENTORY_ENCRYPTION_KEY"]
    assert key not in request("GET", "/api/admin/inventory/status").text


def test_listing_with_wrong_key_returns_clear_error():
    assert upload("import", make_xlsx(sample_rows(2))).status_code == 200
    os.environ["INVENTORY_ENCRYPTION_KEY"] = Fernet.generate_key().decode()
    res = request("GET", f"/api/admin/products/{PID}/inventory")
    assert res.status_code == 503
    assert "tidak cocok" in res.json()["detail"]


def test_header_mismatch_returns_400_with_detail():
    assert upload("import", make_xlsx(sample_rows(2))).status_code == 200
    other = make_xlsx([("a@x.com", "p")], header=("email", "password"))
    res = upload("import", other)
    assert res.status_code == 400
    assert "Header inventory tidak cocok" in res.json()["detail"]


def test_csv_import():
    csv_data = "email,password,recovery_email,2FA\n" + "\n".join(
        f"c{i}@x.com,p{i},r{i}@x.com,f{i}" for i in range(4)
    )
    res = upload("import", csv_data.encode(), filename="data.csv", mime="text/csv")
    assert res.status_code == 200, res.text
    assert res.json()["created"] == 4


def test_transform_inventory_fixed_password_and_domain_only_selected_email_column():
    from fastapi import UploadFile

    data = make_xlsx([
        ("a@old.com", "oldpw1", "recovery@gmail.com", "2fa1"),
        ("b@other.com", "oldpw2", "backup@yahoo.com", "2fa2"),
    ])
    def make_upload():
        return UploadFile(file=io.BytesIO(data), filename="stock.xlsx")

    inspected = run(inventory_transform.inspect_inventory_file(PID, make_upload()))
    assert inspected == {"schema": list(HEADER), "row_count": 2,
                         "email_columns": ["email", "recovery_email"],
                         "password_columns": ["password"]}
    assert "oldpw" not in str(inspected)
    transformed = run(inventory_transform.transform_inventory_file(
        PID, make_upload(), old_domain="old.com", new_domain="new.com",
        email_column="email", password_column="password",
        password_mode="fixed", fixed_password="Fixed!123"))
    assert transformed.headers["x-rows"] == "2"
    assert transformed.headers["x-domain-changes"] == "1"
    assert transformed.headers["x-password-changes"] == "2"
    content = transformed.body
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    rows = list(wb.active.values)
    assert rows[0] == HEADER
    assert rows[1] == ("a@new.com", "Fixed!123", "recovery@gmail.com", "2fa1")
    assert rows[2] == ("b@other.com", "Fixed!123", "backup@yahoo.com", "2fa2")
    assert run(db.inventory_items.count_documents({"product_id": PID})) == 0


def test_transform_inventory_random_password_is_12_chars_and_importable():
    import string
    from fastapi import UploadFile

    data = ("email,kata sandi,recovery_email\n"
            "a@old.com,first,backup@gmail.com\n"
            "b@old.com,second,backup@yahoo.com\n").encode()
    transformed = run(inventory_transform.transform_inventory_file(
        PID, UploadFile(file=io.BytesIO(data), filename="stock.csv"),
        old_domain="", new_domain="@new.com", email_column="email",
        password_column="kata sandi", password_mode="random", fixed_password=""))
    content = transformed.body
    rows = list(openpyxl.load_workbook(io.BytesIO(content), data_only=True).active.values)
    assert rows[1][0] == "a@new.com" and rows[2][0] == "b@new.com"
    assert rows[1][2] == "backup@gmail.com"
    passwords = [row[1] for row in rows[1:]]
    assert passwords[0] != passwords[1]
    for password in passwords:
        assert len(password) == 12
        assert any(char in string.ascii_uppercase for char in password)
        assert any(char in string.ascii_lowercase for char in password)
        assert any(char in string.digits for char in password)
        assert any(char in inventory_transform.PASSWORD_SPECIAL for char in password)
    class InMemoryUpload:
        filename = "transformed.xlsx"

        async def read(self):
            return content

    schema, records = run(admin_routes._parse_inventory_input(InMemoryUpload(), "", {}))
    assert schema == ["email", "kata sandi", "recovery_email"]
    assert len(records) == 2


def test_transform_inventory_rejects_missing_password_header():
    from fastapi import HTTPException, UploadFile

    data = make_xlsx([("a@x.com", "old")], header=("email", "token"))
    with pytest.raises(HTTPException) as exc:
        run(inventory_transform.transform_inventory_file(
            PID, UploadFile(file=io.BytesIO(data), filename="stock.xlsx"),
            old_domain="", new_domain="", email_column="", password_column="token",
            password_mode="random", fixed_password=""))
    assert exc.value.status_code == 400
    assert "password, kata sandi, atau sandi" in exc.value.detail


def test_txt_import_uses_existing_schema():
    assert upload("import", make_xlsx(sample_rows(1, "seed"))).status_code == 200
    txt = "\n".join(f"t{i}@x.com|p{i}|r{i}@x.com|f{i}" for i in range(3))
    res = upload("import", txt.encode(), filename="data.txt", mime="text/plain")
    assert res.status_code == 200, res.text
    assert res.json()["created"] == 3
    assert res.json()["stock"] == 4


def test_xlsx_with_broken_styles_uses_fallback_parser():
    src = io.BytesIO(make_xlsx(sample_rows(5)))
    out = io.BytesIO()
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            payload = zin.read(item.filename)
            if item.filename == "xl/styles.xml":
                payload = b"<styleSheet><broken"
            zout.writestr(item, payload)
    res = upload("import", out.getvalue())
    assert res.status_code == 200, res.text
    assert res.json()["created"] == 5


def test_manual_add_then_duplicate_conflict():
    assert upload("import", make_xlsx(sample_rows(1))).status_code == 200
    payload = {"data": {"email": "m@x.com", "password": "p", "recovery_email": "r@x.com", "2FA": "f"}}
    first = request("POST", f"/api/admin/products/{PID}/inventory/manual", json=payload)
    assert first.status_code == 200, first.text
    assert first.json()["stock"] == 2
    second = request("POST", f"/api/admin/products/{PID}/inventory/manual", json=payload)
    assert second.status_code == 409


def test_import_only_touches_selected_product():
    async def add_other():
        await db.products.insert_one({
            "_id": "other", "name": "Lain", "product_kind": "digital",
            "delivery_type": "inventory", "inventory_enabled": True, "inventory_schema": [],
        })

    run(add_other())
    assert upload("import", make_xlsx(sample_rows(3))).status_code == 200
    assert run(inventory.available_count("other")) == 0
    assert run(inventory.available_count(PID)) == 3


def test_race_duplicate_is_skipped_not_error(monkeypatch):
    rows = [{"email": "a@x.com", "password": "p"}]
    schema = ["email", "password"]
    run(inventory.add_records(PID, rows, schema))

    real_check = inventory._check_records

    async def stale_check(product_id, records, sch):
        # Simulasi request paralel: pengecekan duplikat tidak melihat data yang baru masuk.
        sch, valid, _ = await real_check("tidak-ada", records, sch)
        return sch, valid, []

    monkeypatch.setattr(inventory, "_check_records", stale_check)
    result = run(inventory.add_records(PID, rows, schema))
    assert result["created"] == 0
    assert result["skipped"] == 1


def test_database_write_failure_returns_detail(monkeypatch):
    async def broken_insert_many(*args, **kwargs):
        raise ConnectionError("mongo putus")

    monkeypatch.setattr(AsyncMongoMockCollection, "insert_many", broken_insert_many)
    res = upload("import", make_xlsx(sample_rows(2)))
    assert res.status_code == 500
    assert "ConnectionError" in res.json()["detail"]
    assert "mongo putus" not in res.text  # pesan mentah tidak bocor ke klien


def test_unhandled_exception_is_json_with_detail():
    res = request("GET", "/api/boom")
    assert res.status_code == 500
    assert res.headers["content-type"].startswith("application/json")
    assert "RuntimeError" in res.json()["detail"]
